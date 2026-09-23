import json
import os
import re
import sqlite3
import secrets
from datetime import datetime, timezone, timedelta
from bottle import Bottle, request, response, static_file, run

from app.database import get_db, init_db, hash_password, verify_password
from app.seed import seed_database
from app.gantt_parser import (
    parse_gantt_file, generate_sample_gantt_csv, generate_sample_gantt_excel, AVATAR_COLORS
)
from app.notifier import (
    notify_task_assigned, notify_task_completed, run_all_due_date_checks,
    send_test_email, get_settings, start_background_scheduler
)

app = Bottle()

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
STATIC_DIR = os.path.join(BASE_DIR, "static")

def get_now_iso():
    return datetime.now(timezone.utc).isoformat()

def json_serial_default(obj):
    if isinstance(obj, sqlite3.Row):
        return dict(obj)
    if hasattr(obj, "isoformat"):
        return obj.isoformat()
    try:
        return dict(obj)
    except Exception:
        return str(obj)

def json_response(data, status=200):
    response.status = status
    response.content_type = "application/json"
    return json.dumps(data, default=json_serial_default)

import math

def safe_float(val, default=0.0):
    if val is None or val == "":
        return default
    try:
        f = float(val)
        return default if (math.isnan(f) or math.isinf(f)) else f
    except (ValueError, TypeError):
        return default

def safe_int(val, default=None):
    if val is None or val == "":
        return default
    try:
        return int(val)
    except (ValueError, TypeError):
        return default

def clean_text(val):
    if val is None:
        return ""
    return " ".join(str(val).strip().split())

def record_activity(conn, project_id, user_name, action, details, task_id=None):
    now_str = get_now_iso()
    conn.execute("""
        INSERT INTO activity_logs (project_id, task_id, user_name, action, details, timestamp)
        VALUES (?, ?, ?, ?, ?, ?)
    """, (project_id, task_id, user_name, action, details, now_str))

@app.hook("after_request")
def enable_cors():
    response.headers["Access-Control-Allow-Origin"] = "*"
    response.headers["Access-Control-Allow-Methods"] = "GET, POST, PUT, DELETE, OPTIONS"
    response.headers["Access-Control-Allow-Headers"] = "Origin, Accept, Content-Type, X-Requested-With, X-CSRF-Token, Authorization"

@app.route("/<:re:.*>", method="OPTIONS")
def enable_cors_generic_route():
    return ""

# ==================== STATIC FILES ====================

@app.get("/")
def serve_index():
    response.set_header("Cache-Control", "no-cache, must-revalidate")
    return static_file("index.html", root=STATIC_DIR)

@app.get("/static/<filepath:path>")
def serve_static(filepath):
    response.set_header("Cache-Control", "public, max-age=604800")
    return static_file(filepath, root=STATIC_DIR)

@app.get("/favicon.ico")
def serve_favicon():
    return ""

@app.get("/health")
@app.get("/_health")
@app.get("/api/health")
def health_check():
    return json_response({"status": "healthy", "service": "ProjectPulse", "timestamp": get_now_iso()})

# ==================== AUTHENTICATION ====================

def get_current_user():
    auth_header = request.headers.get("Authorization", "")
    token = None
    if auth_header.startswith("Bearer "):
        token = auth_header[7:].strip()
    if not token:
        token = request.get_cookie("pp_token")
    
    if not token:
        return None
    
    with get_db() as conn:
        session = conn.execute("""
            SELECT s.*, u.id as user_id, u.username, u.email, u.full_name, u.role, u.avatar_color
            FROM sessions s
            JOIN users u ON s.user_id = u.id
            WHERE s.token = ? AND s.expires_at > datetime('now')
        """, (token,)).fetchone()
        return session

def get_all_projects_aggregated(conn):
    return conn.execute("""
        SELECT
            p.*,
            COUNT(t.id) as total_tasks,
            SUM(CASE WHEN t.status = 'done' THEN 1 ELSE 0 END) as completed_tasks,
            SUM(CASE WHEN t.status != 'done' AND t.due_date < date('now') AND t.due_date IS NOT NULL AND t.due_date != '' THEN 1 ELSE 0 END) as overdue_tasks,
            COALESCE(SUM(t.actual_hours), 0) as total_actual_hours,
            COALESCE(SUM(t.estimated_hours), 0) as total_estimated_hours
        FROM projects p
        LEFT JOIN tasks t ON t.project_id = p.id
        GROUP BY p.id
        ORDER BY p.updated_at DESC
    """).fetchall()

def get_bootstrap_payload(conn, user_id=None, active_project_id=None):
    projects = get_all_projects_aggregated(conn)

    if not projects:
        return {
            "projects": [],
            "current_project": None,
            "tasks": []
        }

    target_id = active_project_id
    if not target_id or not any(p["id"] == target_id for p in projects):
        target_id = projects[0]["id"]

    project = conn.execute("SELECT * FROM projects WHERE id = ?", (target_id,)).fetchone()
    members = conn.execute("SELECT * FROM members WHERE project_id = ? ORDER BY name", (target_id,)).fetchall()
    sprints = conn.execute("SELECT * FROM sprints WHERE project_id = ? ORDER BY start_date DESC", (target_id,)).fetchall()
    milestones = conn.execute("SELECT * FROM milestones WHERE project_id = ? ORDER BY due_date ASC", (target_id,)).fetchall()

    current_project = dict(project) if project else None
    if current_project:
        current_project["members"] = members
        current_project["sprints"] = sprints
        current_project["milestones"] = milestones

    tasks_raw = conn.execute("""
        SELECT t.*,
            m.name as assignee_name, m.avatar_color as assignee_avatar, m.role as assignee_role,
            s.name as sprint_name
        FROM tasks t
        LEFT JOIN members m ON t.assignee_id = m.id
        LEFT JOIN sprints s ON t.sprint_id = s.id
        WHERE t.project_id = ?
        ORDER BY t.order_index ASC, t.id ASC
    """, (target_id,)).fetchall()

    subtasks_raw = conn.execute("""
        SELECT s.* FROM subtasks s 
        JOIN tasks t ON s.task_id = t.id 
        WHERE t.project_id = ? 
        ORDER BY s.order_index ASC
    """, (target_id,)).fetchall()

    subtasks_by_task = {}
    for s in subtasks_raw:
        subtasks_by_task.setdefault(s["task_id"], []).append(dict(s))

    timelogs_raw = conn.execute("""
        SELECT tl.task_id, COALESCE(SUM(tl.hours), 0) as total_logged_hours
        FROM timelogs tl
        JOIN tasks t ON tl.task_id = t.id
        WHERE t.project_id = ?
        GROUP BY tl.task_id
    """, (target_id,)).fetchall()
    logged_hours_by_task = {row["task_id"]: row["total_logged_hours"] for row in timelogs_raw}

    tasks = []
    for t in tasks_raw:
        t_dict = dict(t)
        try:
            t_dict["tags"] = json.loads(t_dict["tags"]) if t_dict.get("tags") else []
        except Exception:
            t_dict["tags"] = []
        
        t_subtasks = subtasks_by_task.get(t["id"], [])
        t_dict["subtasks_list"] = t_subtasks
        t_dict["subtask_count"] = len(t_subtasks)
        t_dict["subtask_completed_count"] = sum(1 for s in t_subtasks if s.get("completed"))
        t_dict["logged_hours_sum"] = safe_float(logged_hours_by_task.get(t["id"], 0.0))
        tasks.append(t_dict)

    return {
        "projects": projects,
        "current_project": current_project,
        "tasks": tasks
    }

@app.post("/api/auth/login")
def auth_login():
    data = request.json or {}
    identifier = (data.get("username") or data.get("email") or "").strip().lower()
    password = data.get("password") or ""
    remember = bool(data.get("remember", True))
    active_p_id = data.get("active_project_id")
    active_p_id = int(active_p_id) if active_p_id and str(active_p_id).isdigit() else None

    if not identifier or not password:
        return json_response({"error": "Please enter your username/email and password"}, status=400)

    with get_db() as conn:
        user = conn.execute("""
            SELECT * FROM users
            WHERE LOWER(username) = ? OR LOWER(email) = ?
        """, (identifier, identifier)).fetchone()

        if not user or not verify_password(password, user["password_hash"]):
            return json_response({"error": "Invalid username or password"}, status=401)

        token = secrets.token_hex(32)
        now_dt = datetime.now(timezone.utc)
        now_str = now_dt.isoformat()
        expires_dt = now_dt + timedelta(days=30 if remember else 1)
        expires_str = expires_dt.isoformat()

        conn.execute("""
            INSERT INTO sessions (token, user_id, created_at, expires_at)
            VALUES (?, ?, ?, ?)
        """, (token, user["id"], now_str, expires_str))

        conn.execute("UPDATE users SET last_login = ? WHERE id = ?", (now_str, user["id"]))

        user_dict = {
            "id": user["id"],
            "username": user["username"],
            "email": user["email"],
            "full_name": user["full_name"],
            "role": user["role"],
            "avatar_color": user["avatar_color"]
        }

        # Precompute full bootstrap payload for instant 0ms landing
        bootstrap = get_bootstrap_payload(conn, user["id"], active_p_id)

        # Set session cookie
        max_age = 30 * 86400 if remember else 86400
        response.set_cookie("pp_token", token, path="/", max_age=max_age, httponly=False, samesite="Lax")

        return json_response({
            "success": True,
            "token": token,
            "user": user_dict,
            "projects": bootstrap["projects"],
            "current_project": bootstrap["current_project"],
            "tasks": bootstrap["tasks"]
        })

@app.post("/api/auth/register")
def auth_register():
    data = request.json or {}
    full_name = (data.get("full_name") or "").strip()
    username = (data.get("username") or "").strip().lower()
    email = (data.get("email") or "").strip().lower()
    password = data.get("password") or ""

    if not full_name:
        return json_response({"error": "Full name is required"}, status=400)
    if not username or len(username) < 3:
        return json_response({"error": "Username must be at least 3 characters"}, status=400)
    if not email or "@" not in email:
        return json_response({"error": "Please provide a valid email address"}, status=400)
    if not password or len(password) < 6:
        return json_response({"error": "Password must be at least 6 characters"}, status=400)

    now_dt = datetime.now(timezone.utc)
    now_str = now_dt.isoformat()
    pwd_hash = hash_password(password)

    avatar_colors = ["#3B82F6", "#8B5CF6", "#EC4899", "#10B981", "#F59E0B", "#06B6D4", "#6366F1"]
    avatar_color = avatar_colors[len(username) % len(avatar_colors)]

    with get_db() as conn:
        existing = conn.execute("""
            SELECT username, email FROM users
            WHERE LOWER(username) = ? OR LOWER(email) = ?
        """, (username, email)).fetchone()

        if existing:
            if existing["username"].lower() == username:
                return json_response({"error": "This username is already registered"}, status=400)
            else:
                return json_response({"error": "An account with this email already exists"}, status=400)

        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO users (username, email, password_hash, full_name, role, avatar_color, created_at, last_login)
            VALUES (?, ?, ?, ?, 'manager', ?, ?, ?)
        """, (username, email, pwd_hash, full_name, avatar_color, now_str, now_str))
        user_id = cursor.lastrowid

        token = secrets.token_hex(32)
        expires_str = (now_dt + timedelta(days=30)).isoformat()
        cursor.execute("""
            INSERT INTO sessions (token, user_id, created_at, expires_at)
            VALUES (?, ?, ?, ?)
        """, (token, user_id, now_str, expires_str))

        user_dict = {
            "id": user_id,
            "username": username,
            "email": email,
            "full_name": full_name,
            "role": "manager",
            "avatar_color": avatar_color
        }

        bootstrap = get_bootstrap_payload(conn, user_id)
        response.set_cookie("pp_token", token, path="/", max_age=30*86400, httponly=False, samesite="Lax")

        return json_response({
            "success": True,
            "token": token,
            "user": user_dict,
            "projects": bootstrap["projects"],
            "current_project": bootstrap["current_project"],
            "tasks": bootstrap["tasks"]
        }, status=201)

@app.get("/api/auth/me")
def auth_me():
    user = get_current_user()
    if not user:
        return json_response({"authenticated": False}, status=401)
    
    active_p_id = request.query.get("active_project_id")
    active_p_id = int(active_p_id) if active_p_id and str(active_p_id).isdigit() else None

    with get_db() as conn:
        bootstrap = get_bootstrap_payload(conn, user["user_id"], active_p_id)

    return json_response({
        "authenticated": True,
        "user": {
            "id": user["user_id"],
            "username": user["username"],
            "email": user["email"],
            "full_name": user["full_name"],
            "role": user["role"],
            "avatar_color": user["avatar_color"]
        },
        "projects": bootstrap["projects"],
        "current_project": bootstrap["current_project"],
        "tasks": bootstrap["tasks"]
    })

@app.get("/api/bootstrap")
def api_bootstrap():
    user = get_current_user()
    active_p_id = request.query.get("active_project_id")
    active_p_id = int(active_p_id) if active_p_id and str(active_p_id).isdigit() else None
    with get_db() as conn:
        bootstrap = get_bootstrap_payload(conn, user["user_id"] if user else None, active_p_id)
        return json_response(bootstrap)

@app.post("/api/auth/logout")
def auth_logout():
    auth_header = request.headers.get("Authorization", "")
    token = None
    if auth_header.startswith("Bearer "):
        token = auth_header[7:].strip()
    if not token:
        token = request.get_cookie("pp_token")
    
    if token:
        with get_db() as conn:
            conn.execute("DELETE FROM sessions WHERE token = ?", (token,))
    
    response.delete_cookie("pp_token", path="/")
    return json_response({"success": True})

# ==================== PROJECTS ====================

@app.get("/api/projects")
def get_projects():
    with get_db() as conn:
        projects = get_all_projects_aggregated(conn)
        return json_response(projects)

@app.post("/api/projects")
def create_project():
    data = request.json or {}
    name = data.get("name", "").strip()
    if not name:
        return json_response({"error": "Project name is required"}, status=400)
    
    desc = data.get("description", "")
    color = data.get("color", "#3B82F6")
    now_str = get_now_iso()

    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO projects (name, description, color, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?)
        """, (name, desc, color, now_str, now_str))
        p_id = cursor.lastrowid

        cursor.execute("""
            INSERT INTO members (project_id, name, email, role, avatar_color)
            VALUES (?, ?, ?, ?, ?)
        """, (p_id, "Project Lead", "lead@company.internal", "Owner", color))

        record_activity(conn, p_id, "System", "Project Created", f'Created project "{name}"')
        project = conn.execute("SELECT * FROM projects WHERE id = ?", (p_id,)).fetchone()
        members = conn.execute("SELECT * FROM members WHERE project_id = ?", (p_id,)).fetchall()
        
        proj_dict = dict(project)
        proj_dict["members"] = members
        proj_dict["sprints"] = []
        proj_dict["milestones"] = []
        proj_dict["total_tasks"] = 0
        proj_dict["completed_tasks"] = 0
        proj_dict["overdue_tasks"] = 0
        proj_dict["total_actual_hours"] = 0
        proj_dict["total_estimated_hours"] = 0
        return json_response(proj_dict, status=201)

def deduplicate_project_members(conn, project_id):
    """
    Ensures no duplicate members with matching normalized names exist per project.
    Reassigns any tasks from duplicate IDs to the canonical ID, then removes duplicate records.
    """
    rows = conn.execute(
        "SELECT id, name FROM members WHERE project_id = ? ORDER BY id ASC",
        (project_id,)
    ).fetchall()
    seen = {}
    for r in rows:
        clean_name = " ".join(r["name"].strip().split()).lower()
        if not clean_name:
            continue
        if clean_name in seen:
            canonical_id = seen[clean_name]
            dup_id = r["id"]
            conn.execute(
                "UPDATE tasks SET assignee_id = ? WHERE project_id = ? AND assignee_id = ?",
                (canonical_id, project_id, dup_id)
            )
            conn.execute("DELETE FROM members WHERE id = ?", (dup_id,))
        else:
            seen[clean_name] = r["id"]

@app.get("/api/projects/<project_id:int>")
def get_project(project_id):
    with get_db() as conn:
        project = conn.execute("SELECT * FROM projects WHERE id = ?", (project_id,)).fetchone()
        if not project:
            return json_response({"error": "Project not found"}, status=404)
        
        deduplicate_project_members(conn, project_id)
        members = conn.execute("SELECT * FROM members WHERE project_id = ? ORDER BY name", (project_id,)).fetchall()
        sprints = conn.execute("SELECT * FROM sprints WHERE project_id = ? ORDER BY start_date DESC", (project_id,)).fetchall()
        milestones = conn.execute("SELECT * FROM milestones WHERE project_id = ? ORDER BY due_date ASC", (project_id,)).fetchall()
        
        result = dict(project)
        result["members"] = members
        result["sprints"] = sprints
        result["milestones"] = milestones
        return json_response(result)

@app.put("/api/projects/<project_id:int>")
def update_project(project_id):
    data = request.json or {}
    with get_db() as conn:
        project = conn.execute("SELECT * FROM projects WHERE id = ?", (project_id,)).fetchone()
        if not project:
            return json_response({"error": "Project not found"}, status=404)
        
        name = data.get("name", project["name"])
        desc = data.get("description", project["description"])
        color = data.get("color", project["color"])
        now_str = get_now_iso()

        conn.execute("""
            UPDATE projects SET name = ?, description = ?, color = ?, updated_at = ?
            WHERE id = ?
        """, (name, desc, color, now_str, project_id))
        
        record_activity(conn, project_id, "User", "Project Updated", "Updated project settings")
        updated = conn.execute("SELECT * FROM projects WHERE id = ?", (project_id,)).fetchone()
        return json_response(updated)

@app.delete("/api/projects/<project_id:int>")
def delete_project(project_id):
    with get_db() as conn:
        project = conn.execute("SELECT id FROM projects WHERE id = ?", (project_id,)).fetchone()
        if not project:
            return json_response({"error": "Project not found"}, status=404)
        conn.execute("DELETE FROM projects WHERE id = ?", (project_id,))
        return json_response({
            "success": True,
            "message": "Project deleted",
            "deleted_id": project_id
        })

# ==================== MEMBERS ====================

@app.get("/api/projects/<project_id:int>/members")
def get_members(project_id):
    with get_db() as conn:
        deduplicate_project_members(conn, project_id)
        members = conn.execute("SELECT * FROM members WHERE project_id = ? ORDER BY name", (project_id,)).fetchall()
        return json_response(members)

@app.post("/api/projects/<project_id:int>/members")
def add_member(project_id):
    data = request.json or {}
    name = " ".join(data.get("name", "").strip().split())
    if not name:
        return json_response({"error": "Name is required"}, status=400)
    
    email = data.get("email", "").strip()
    role = data.get("role", "Member")
    avatar_color = data.get("avatar_color", "#6366F1")

    with get_db() as conn:
        cursor = conn.cursor()
        existing = conn.execute(
            "SELECT * FROM members WHERE project_id = ? AND LOWER(TRIM(name)) = LOWER(TRIM(?))",
            (project_id, name)
        ).fetchone()
        if existing:
            return json_response(dict(existing), status=200)

        cursor.execute("""
            INSERT INTO members (project_id, name, email, role, avatar_color)
            VALUES (?, ?, ?, ?, ?)
        """, (project_id, name, email, role, avatar_color))
        m_id = cursor.lastrowid
        record_activity(conn, project_id, "Admin", "Member Added", f'Added member "{name}" ({role})')
        member = conn.execute("SELECT * FROM members WHERE id = ?", (m_id,)).fetchone()
        return json_response(dict(member), status=201)

@app.delete("/api/projects/<project_id:int>/members/<member_id:int>")
def delete_member(project_id, member_id):
    with get_db() as conn:
        member = conn.execute("SELECT name FROM members WHERE id = ? AND project_id = ?", (member_id, project_id)).fetchone()
        if member:
            target_name = member["name"].strip()
            # Clean up all duplicate records with same name (case-insensitive & trimmed)
            matched_members = conn.execute(
                "SELECT id FROM members WHERE project_id = ? AND (id = ? OR LOWER(TRIM(name)) = LOWER(TRIM(?)))",
                (project_id, member_id, target_name)
            ).fetchall()
            all_ids = [m["id"] for m in matched_members] or [member_id]
            
            placeholders = ",".join("?" * len(all_ids))
            conn.execute(f"UPDATE tasks SET assignee_id = NULL WHERE project_id = ? AND assignee_id IN ({placeholders})", [project_id] + all_ids)
            conn.execute(f"DELETE FROM members WHERE project_id = ? AND id IN ({placeholders})", [project_id] + all_ids)
            record_activity(conn, project_id, "Admin", "Member Removed", f'Removed member "{target_name}"')
            return json_response({"success": True, "deleted_member_id": member_id, "deleted_ids": all_ids, "deleted_name": target_name})
        else:
            conn.execute("UPDATE tasks SET assignee_id = NULL WHERE project_id = ? AND assignee_id = ?", (project_id, member_id))
            conn.execute("DELETE FROM members WHERE id = ? AND project_id = ?", (member_id, project_id))
            return json_response({"success": True, "deleted_member_id": member_id})

@app.delete("/api/members/<member_id:int>")
def delete_member_direct(member_id):
    with get_db() as conn:
        member = conn.execute("SELECT project_id, name FROM members WHERE id = ?", (member_id,)).fetchone()
        if not member:
            return json_response({"error": "Member not found"}, status=404)
        project_id = member["project_id"]
        target_name = member["name"].strip()
        matched_members = conn.execute(
            "SELECT id FROM members WHERE project_id = ? AND (id = ? OR LOWER(TRIM(name)) = LOWER(TRIM(?)))",
            (project_id, member_id, target_name)
        ).fetchall()
        all_ids = [m["id"] for m in matched_members] or [member_id]
        
        placeholders = ",".join("?" * len(all_ids))
        conn.execute(f"UPDATE tasks SET assignee_id = NULL WHERE project_id = ? AND assignee_id IN ({placeholders})", [project_id] + all_ids)
        conn.execute(f"DELETE FROM members WHERE project_id = ? AND id IN ({placeholders})", [project_id] + all_ids)
        record_activity(conn, project_id, "Admin", "Member Removed", f'Removed member "{target_name}"')
        return json_response({"success": True, "deleted_member_id": member_id, "deleted_ids": all_ids, "deleted_name": target_name})

# ==================== SPRINTS ====================

@app.get("/api/projects/<project_id:int>/sprints")
def get_sprints(project_id):
    with get_db() as conn:
        sprints = conn.execute("""
            SELECT s.*,
                (SELECT COUNT(*) FROM tasks WHERE sprint_id = s.id) as total_tasks,
                (SELECT COUNT(*) FROM tasks WHERE sprint_id = s.id AND status = 'done') as completed_tasks,
                (SELECT COALESCE(SUM(estimated_hours), 0) FROM tasks WHERE sprint_id = s.id) as total_estimated_hours,
                (SELECT COALESCE(SUM(actual_hours), 0) FROM tasks WHERE sprint_id = s.id) as total_actual_hours
            FROM sprints s
            WHERE s.project_id = ?
            ORDER BY s.start_date DESC
        """, (project_id,)).fetchall()
        return json_response(sprints)

@app.post("/api/projects/<project_id:int>/sprints")
def create_sprint(project_id):
    data = request.json or {}
    name = data.get("name", "").strip()
    if not name:
        return json_response({"error": "Sprint name is required"}, status=400)
    
    goal = data.get("goal", "")
    start_date = data.get("start_date")
    end_date = data.get("end_date")
    status = data.get("status", "planning")

    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO sprints (project_id, name, goal, start_date, end_date, status)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (project_id, name, goal, start_date, end_date, status))
        s_id = cursor.lastrowid
        record_activity(conn, project_id, "Manager", "Sprint Created", f'Created sprint "{name}"')
        sprint = conn.execute("SELECT * FROM sprints WHERE id = ?", (s_id,)).fetchone()
        return json_response(sprint, status=201)

@app.put("/api/projects/<project_id:int>/sprints/<sprint_id:int>")
def update_sprint(project_id, sprint_id):
    data = request.json or {}
    with get_db() as conn:
        sprint = conn.execute("SELECT * FROM sprints WHERE id = ? AND project_id = ?", (sprint_id, project_id)).fetchone()
        if not sprint:
            return json_response({"error": "Sprint not found"}, status=404)
        
        name = data.get("name", sprint["name"])
        goal = data.get("goal", sprint["goal"])
        start_date = data.get("start_date", sprint["start_date"])
        end_date = data.get("end_date", sprint["end_date"])
        status = data.get("status", sprint["status"])

        if status == "active" and sprint["status"] != "active":
            conn.execute("UPDATE sprints SET status = 'planning' WHERE project_id = ? AND status = 'active'", (project_id,))

        conn.execute("""
            UPDATE sprints SET name = ?, goal = ?, start_date = ?, end_date = ?, status = ?
            WHERE id = ?
        """, (name, goal, start_date, end_date, status, sprint_id))

        record_activity(conn, project_id, "Manager", "Sprint Updated", f'Updated sprint "{name}" ({status})')
        updated = conn.execute("SELECT * FROM sprints WHERE id = ?", (sprint_id,)).fetchone()
        return json_response(updated)

@app.delete("/api/projects/<project_id:int>/sprints/<sprint_id:int>")
def delete_sprint(project_id, sprint_id):
    with get_db() as conn:
        conn.execute("DELETE FROM sprints WHERE id = ? AND project_id = ?", (sprint_id, project_id))
        return json_response({"success": True})

# ==================== MILESTONES ====================

@app.get("/api/projects/<project_id:int>/milestones")
def get_milestones(project_id):
    with get_db() as conn:
        milestones = conn.execute("SELECT * FROM milestones WHERE project_id = ? ORDER BY due_date ASC", (project_id,)).fetchall()
        return json_response(milestones)

@app.post("/api/projects/<project_id:int>/milestones")
def create_milestone(project_id):
    data = request.json or {}
    title = data.get("title", "").strip()
    due_date = data.get("due_date", "").strip()
    if not title or not due_date:
        return json_response({"error": "Title and due date are required"}, status=400)
    
    status = data.get("status", "pending")
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO milestones (project_id, title, due_date, status)
            VALUES (?, ?, ?, ?)
        """, (project_id, title, due_date, status))
        m_id = cursor.lastrowid
        record_activity(conn, project_id, "Lead", "Milestone Added", f'Added milestone "{title}"')
        m = conn.execute("SELECT * FROM milestones WHERE id = ?", (m_id,)).fetchone()
        return json_response(m, status=201)

@app.put("/api/projects/<project_id:int>/milestones/<milestone_id:int>")
def update_milestone(project_id, milestone_id):
    data = request.json or {}
    with get_db() as conn:
        m = conn.execute("SELECT * FROM milestones WHERE id = ? AND project_id = ?", (milestone_id, project_id)).fetchone()
        if not m:
            return json_response({"error": "Milestone not found"}, status=404)
        
        title = data.get("title", m["title"])
        due_date = data.get("due_date", m["due_date"])
        status = data.get("status", m["status"])

        conn.execute("UPDATE milestones SET title = ?, due_date = ?, status = ? WHERE id = ?", (title, due_date, status, milestone_id))
        updated = conn.execute("SELECT * FROM milestones WHERE id = ?", (milestone_id,)).fetchone()
        return json_response(updated)

@app.delete("/api/projects/<project_id:int>/milestones/<milestone_id:int>")
def delete_milestone(project_id, milestone_id):
    with get_db() as conn:
        conn.execute("DELETE FROM milestones WHERE id = ? AND project_id = ?", (milestone_id, project_id))
        return json_response({"success": True})

# ==================== TASKS ====================

@app.get("/api/projects/<project_id:int>/tasks")
def get_tasks(project_id):
    sprint_id = request.query.get("sprint_id")
    status = request.query.get("status")
    priority = request.query.get("priority")
    assignee_id = request.query.get("assignee_id")
    search = request.query.get("search")

    with get_db() as conn:
        query = """
            SELECT t.*,
                m.name as assignee_name, m.avatar_color as assignee_avatar, m.role as assignee_role,
                s.name as sprint_name
            FROM tasks t
            LEFT JOIN members m ON t.assignee_id = m.id
            LEFT JOIN sprints s ON t.sprint_id = s.id
            WHERE t.project_id = ?
        """
        params = [project_id]

        if sprint_id:
            if sprint_id == "backlog":
                query += " AND t.sprint_id IS NULL"
            elif sprint_id.isdigit():
                query += " AND t.sprint_id = ?"
                params.append(int(sprint_id))

        if status:
            query += " AND t.status = ?"
            params.append(status)

        if priority:
            query += " AND t.priority = ?"
            params.append(priority)

        if assignee_id and assignee_id.isdigit():
            query += " AND t.assignee_id = ?"
            params.append(int(assignee_id))

        if search:
            query += " AND (t.title LIKE ? OR t.description LIKE ? OR t.tags LIKE ?)"
            s_param = f"%{search}%"
            params.extend([s_param, s_param, s_param])

        query += " ORDER BY t.order_index ASC, t.id ASC"
        tasks = conn.execute(query, params).fetchall()

        # Batch fetch all subtasks for this project in 1 single query (eliminates N+1 latency)
        subtasks_raw = conn.execute("""
            SELECT s.* FROM subtasks s 
            JOIN tasks t ON s.task_id = t.id 
            WHERE t.project_id = ? 
            ORDER BY s.order_index ASC
        """, (project_id,)).fetchall()
        
        subtasks_by_task = {}
        for s in subtasks_raw:
            subtasks_by_task.setdefault(s["task_id"], []).append(dict(s))

        # Batch fetch all timelogs in 1 single query
        timelogs_raw = conn.execute("""
            SELECT tl.task_id, COALESCE(SUM(tl.hours), 0) as total_logged_hours
            FROM timelogs tl
            JOIN tasks t ON tl.task_id = t.id
            WHERE t.project_id = ?
            GROUP BY tl.task_id
        """, (project_id,)).fetchall()
        logged_hours_by_task = {row["task_id"]: row["total_logged_hours"] for row in timelogs_raw}

        # Batch fetch all task resources in 1 single query
        resources_raw = conn.execute("""
            SELECT tr.id as task_resource_id, r.id, tr.resource_id, tr.task_id, tr.role, tr.responsibility,
                   r.resource_code, r.name, r.name as resource_name, r.type, r.type as resource_type,
                   r.category, r.category as resource_category, r.department, r.department as resource_department
            FROM task_resources tr
            JOIN resources r ON tr.resource_id = r.id
            WHERE tr.project_id = ?
            ORDER BY r.name ASC
        """, (project_id,)).fetchall()
        resources_by_task = {}
        for r in resources_raw:
            resources_by_task.setdefault(r["task_id"], []).append(dict(r))

        result = []
        for t in tasks:
            t_dict = dict(t)
            try:
                t_dict["tags"] = json.loads(t_dict["tags"]) if t_dict["tags"] else []
            except Exception:
                t_dict["tags"] = []
            
            t_subtasks = subtasks_by_task.get(t["id"], [])
            t_dict["subtasks_list"] = t_subtasks
            t_dict["subtask_count"] = len(t_subtasks)
            t_dict["subtask_completed_count"] = sum(1 for s in t_subtasks if s.get("completed"))
            t_dict["logged_hours_sum"] = safe_float(logged_hours_by_task.get(t["id"], 0.0))
            t_dict["resources"] = resources_by_task.get(t["id"], [])
            result.append(t_dict)

        return json_response(result)

@app.post("/api/projects/<project_id:int>/tasks")
def create_task(project_id):
    try:
        data = request.json or {}
    except Exception:
        data = {}
        
    title = clean_text(data.get("title"))
    if not title:
        return json_response({"error": "Task title is required"}, status=400)
    
    desc = str(data.get("description") or "")
    status = str(data.get("status") or "todo")
    priority = str(data.get("priority") or "medium")
    
    sprint_val = data.get("sprint_id")
    sprint_id = safe_int(sprint_val)
    if sprint_id is not None and sprint_id <= 0:
        sprint_id = None
        
    start_date = data.get("start_date")
    start_date = str(start_date).strip() if start_date and str(start_date).strip() not in ("", "null", "undefined", "None") else None
    
    due_date = data.get("due_date")
    due_date = str(due_date).strip() if due_date and str(due_date).strip() not in ("", "null", "undefined", "None") else None
    
    est_hours = safe_float(data.get("estimated_hours"), 0.0)
    act_hours = safe_float(data.get("actual_hours"), 0.0)
    
    raw_tags = data.get("tags", [])
    if isinstance(raw_tags, list):
        tags_list = [str(t).strip().lstrip('#') for t in raw_tags if str(t).strip()]
    elif isinstance(raw_tags, str):
        tags_list = [t.strip().lstrip('#') for t in raw_tags.split(',') if t.strip()]
    else:
        tags_list = []
    tags_json = json.dumps(tags_list)
    
    subtasks = data.get("subtasks", [])
    now_str = get_now_iso()

    position = data.get("position")  # "end", "start", "after_<id>", "before_<id>"
    insert_after_id = safe_int(data.get("insert_after_id"))
    insert_before_id = safe_int(data.get("insert_before_id"))
    custom_order_index = safe_int(data.get("order_index"))

    if position and isinstance(position, str):
        if position.startswith("after_"):
            try:
                insert_after_id = int(position.split("_")[1])
            except ValueError:
                pass
        elif position.startswith("before_"):
            try:
                insert_before_id = int(position.split("_")[1])
            except ValueError:
                pass

    with get_db() as conn:
        cursor = conn.cursor()

        # Handle assignee resolution with high reliability
        assignee_id = None
        assignee_name_input = clean_text(data.get("assignee_name") or data.get("new_assignee_name"))
        raw_assignee_id = data.get("assignee_id")
        
        if assignee_name_input:
            existing_m = conn.execute(
                "SELECT id FROM members WHERE project_id = ? AND LOWER(TRIM(name)) = LOWER(TRIM(?)) ORDER BY id ASC LIMIT 1",
                (project_id, assignee_name_input)
            ).fetchone()
            if existing_m:
                assignee_id = existing_m["id"]
            else:
                avatar_colors = ["#3B82F6", "#6366F1", "#8B5CF6", "#EC4899", "#10B981", "#F59E0B", "#14B8A6", "#F97316"]
                import random
                color = random.choice(avatar_colors)
                cursor.execute("""
                    INSERT INTO members (project_id, name, email, role, avatar_color)
                    VALUES (?, ?, ?, ?, ?)
                """, (project_id, assignee_name_input, "", "Member", color))
                assignee_id = cursor.lastrowid
                record_activity(conn, project_id, "User", "Member Added", f'Added member "{assignee_name_input}"')
        elif raw_assignee_id is not None:
            parsed_aid = safe_int(raw_assignee_id)
            if parsed_aid and parsed_aid > 0:
                m_check = conn.execute("SELECT id FROM members WHERE id = ? AND project_id = ?", (parsed_aid, project_id)).fetchone()
                if m_check:
                    assignee_id = m_check["id"]
        
        target_order = None
        if insert_after_id:
            pred = conn.execute("SELECT order_index FROM tasks WHERE id = ? AND project_id = ?", (insert_after_id, project_id)).fetchone()
            if pred:
                target_order = pred["order_index"] + 1
                conn.execute(
                    "UPDATE tasks SET order_index = order_index + 1 WHERE project_id = ? AND order_index >= ?",
                    (project_id, target_order)
                )
        elif insert_before_id:
            succ = conn.execute("SELECT order_index FROM tasks WHERE id = ? AND project_id = ?", (insert_before_id, project_id)).fetchone()
            if succ:
                target_order = succ["order_index"]
                conn.execute(
                    "UPDATE tasks SET order_index = order_index + 1 WHERE project_id = ? AND order_index >= ?",
                    (project_id, target_order)
                )
        elif position == "start":
            first_t = conn.execute("SELECT order_index FROM tasks WHERE project_id = ? ORDER BY order_index ASC, id ASC LIMIT 1", (project_id,)).fetchone()
            if first_t:
                target_order = first_t["order_index"]
                conn.execute(
                    "UPDATE tasks SET order_index = order_index + 1 WHERE project_id = ? AND order_index >= ?",
                    (project_id, target_order)
                )
            else:
                target_order = 0
        elif custom_order_index is not None:
            target_order = custom_order_index
            conn.execute(
                "UPDATE tasks SET order_index = order_index + 1 WHERE project_id = ? AND order_index >= ?",
                (project_id, target_order)
            )

        if target_order is None:
            max_order = conn.execute(
                "SELECT COALESCE(MAX(order_index), -1) as max_idx FROM tasks WHERE project_id = ?",
                (project_id,)
            ).fetchone()["max_idx"]
            target_order = max_order + 1

        cursor.execute("""
            INSERT INTO tasks (
                project_id, sprint_id, title, description, status, priority,
                order_index, start_date, due_date, estimated_hours, actual_hours,
                assignee_id, tags, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            project_id, sprint_id, title, desc, status, priority,
            target_order, start_date, due_date, est_hours, act_hours,
            assignee_id, tags_json, now_str, now_str
        ))
        t_id = cursor.lastrowid

        if isinstance(subtasks, list):
            for idx, sub_title in enumerate(subtasks):
                if isinstance(sub_title, str) and sub_title.strip():
                    cursor.execute("""
                        INSERT INTO subtasks (task_id, title, completed, order_index)
                        VALUES (?, ?, 0, ?)
                    """, (t_id, sub_title.strip(), idx))

        # Mapped Resources if provided
        raw_res_ids = data.get("resource_ids")
        if isinstance(raw_res_ids, list):
            for r_id in raw_res_ids:
                try:
                    r_int = int(r_id)
                    cursor.execute("""
                        INSERT OR IGNORE INTO task_resources (task_id, resource_id, project_id, created_at)
                        VALUES (?, ?, ?, ?)
                    """, (t_id, r_int, project_id, now_str))
                except Exception:
                    pass

        record_activity(conn, project_id, "User", "Task Created", f'Created task "{title}"', task_id=t_id)

        if assignee_id:
            try:
                notify_task_assigned(conn, t_id)
            except Exception as e:
                print(f"[Notifier] Error sending assignment notification: {e}")

        task_res = get_task_dict(conn, t_id)
        return json_response(task_res)

def get_task_dict(conn, task_id: int):
    task = conn.execute("""
        SELECT t.*,
            m.name as assignee_name, m.avatar_color as assignee_avatar, m.role as assignee_role,
            s.name as sprint_name,
            p.name as project_name, p.color as project_color
        FROM tasks t
        LEFT JOIN members m ON t.assignee_id = m.id
        LEFT JOIN sprints s ON t.sprint_id = s.id
        LEFT JOIN projects p ON t.project_id = p.id
        WHERE t.id = ?
    """, (task_id,)).fetchone()

    if not task:
        return None

    t_dict = dict(task)
    try:
        t_dict["tags"] = json.loads(t_dict["tags"]) if t_dict["tags"] else []
    except Exception:
        t_dict["tags"] = []

    subtasks_list = [dict(s) for s in conn.execute("SELECT * FROM subtasks WHERE task_id = ? ORDER BY order_index ASC", (task_id,)).fetchall()]
    timelogs_list = [dict(tl) for tl in conn.execute("""
        SELECT tl.*, m.name as member_name, m.avatar_color as member_avatar
        FROM timelogs tl
        LEFT JOIN members m ON tl.member_id = m.id
        WHERE tl.task_id = ?
        ORDER BY tl.logged_date DESC, tl.id DESC
    """, (task_id,)).fetchall()]

    task_resources_list = [dict(tr) for tr in conn.execute("""
        SELECT tr.id as task_resource_id, r.id, tr.resource_id, tr.task_id, tr.role, tr.responsibility,
               r.resource_code, r.name, r.name as resource_name, r.type, r.type as resource_type,
               r.category, r.category as resource_category, r.department, r.department as resource_department
        FROM task_resources tr
        JOIN resources r ON tr.resource_id = r.id
        WHERE tr.task_id = ?
        ORDER BY r.name ASC
    """, (task_id,)).fetchall()]

    t_dict["subtasks"] = subtasks_list
    t_dict["subtasks_list"] = subtasks_list
    t_dict["subtask_count"] = len(subtasks_list)
    t_dict["subtask_completed_count"] = sum(1 for s in subtasks_list if s.get("completed") == 1)
    t_dict["timelogs"] = timelogs_list
    t_dict["logged_hours_sum"] = sum(safe_float(tl.get("hours"), 0.0) for tl in timelogs_list)
    t_dict["resources"] = task_resources_list
    t_dict["activities"] = [dict(a) for a in conn.execute("SELECT * FROM activity_logs WHERE task_id = ? ORDER BY timestamp DESC LIMIT 20", (task_id,)).fetchall()]

    return t_dict

@app.get("/api/tasks/<task_id:int>")
def get_task(task_id):
    with get_db() as conn:
        task_dict = get_task_dict(conn, task_id)
        if not task_dict:
            return json_response({"error": "Task not found"}, status=404)
        return json_response(task_dict)

@app.put("/api/tasks/<task_id:int>")
def update_task(task_id):
    try:
        data = request.json or {}
    except Exception:
        data = {}
        
    with get_db() as conn:
        task = conn.execute("SELECT * FROM tasks WHERE id = ?", (task_id,)).fetchone()
        if not task:
            return json_response({"error": "Task not found"}, status=404)

        project_id = task["project_id"]
        cursor = conn.cursor()
        
        # 1. Text & Enum fields
        title = clean_text(data["title"]) if "title" in data and clean_text(data["title"]) else task["title"]
        desc = str(data["description"]) if "description" in data and data["description"] is not None else task["description"]
        status = str(data["status"]) if "status" in data and data["status"] else task["status"]
        priority = str(data["priority"]) if "priority" in data and data["priority"] else task["priority"]
        
        # 2. Sprint ID
        if "sprint_id" in data:
            sprint_val = data["sprint_id"]
            parsed_sprint = safe_int(sprint_val)
            sprint_id = parsed_sprint if (parsed_sprint and parsed_sprint > 0) else None
        else:
            sprint_id = task["sprint_id"]

        # 3. Assignee resolution (handles custom name input, dropdown id, and unassignment)
        assignee_id = task["assignee_id"]
        assignee_name_input = clean_text(data.get("assignee_name") or data.get("new_assignee_name"))
        
        if assignee_name_input:
            existing_m = conn.execute(
                "SELECT id FROM members WHERE project_id = ? AND LOWER(TRIM(name)) = LOWER(TRIM(?)) ORDER BY id ASC LIMIT 1",
                (project_id, assignee_name_input)
            ).fetchone()
            if existing_m:
                assignee_id = existing_m["id"]
            else:
                avatar_colors = ["#3B82F6", "#6366F1", "#8B5CF6", "#EC4899", "#10B981", "#F59E0B", "#14B8A6", "#F97316"]
                import random
                color = random.choice(avatar_colors)
                cursor.execute("""
                    INSERT INTO members (project_id, name, email, role, avatar_color)
                    VALUES (?, ?, ?, ?, ?)
                """, (project_id, assignee_name_input, "", "Member", color))
                assignee_id = cursor.lastrowid
                record_activity(conn, project_id, "User", "Member Added", f'Added member "{assignee_name_input}"')
        elif "assignee_id" in data:
            assignee_val = data["assignee_id"]
            parsed_aid = safe_int(assignee_val)
            if parsed_aid and parsed_aid > 0:
                m_check = conn.execute("SELECT id FROM members WHERE id = ? AND project_id = ?", (parsed_aid, project_id)).fetchone()
                if m_check:
                    assignee_id = m_check["id"]
                else:
                    assignee_id = None
            else:
                assignee_id = None
        
        # 4. Dates
        if "start_date" in data:
            sd = data["start_date"]
            start_date = str(sd).strip() if sd and str(sd).strip() not in ("", "null", "undefined", "None") else None
        else:
            start_date = task["start_date"]

        if "due_date" in data:
            dd = data["due_date"]
            due_date = str(dd).strip() if dd and str(dd).strip() not in ("", "null", "undefined", "None") else None
        else:
            due_date = task["due_date"]

        # 5. Hours & Order Index
        if "order_index" in data:
            order_index = safe_int(data["order_index"], task["order_index"])
        else:
            order_index = task["order_index"]

        if "estimated_hours" in data:
            est_hours = safe_float(data["estimated_hours"], safe_float(task["estimated_hours"], 0.0))
        else:
            est_hours = safe_float(task["estimated_hours"], 0.0)

        if "actual_hours" in data:
            act_hours = safe_float(data["actual_hours"], safe_float(task["actual_hours"], 0.0))
        else:
            act_hours = safe_float(task["actual_hours"], 0.0)

        # 6. Tags
        if "tags" in data:
            raw_tags = data["tags"]
            if isinstance(raw_tags, list):
                tags_list = [str(t).strip().lstrip('#') for t in raw_tags if str(t).strip()]
            elif isinstance(raw_tags, str):
                tags_list = [t.strip().lstrip('#') for t in raw_tags.split(',') if t.strip()]
            else:
                tags_list = []
            tags_json = json.dumps(tags_list)
        else:
            tags_json = task["tags"] or "[]"
            
        now_str = get_now_iso()

        conn.execute("""
            UPDATE tasks SET
                title = ?, description = ?, status = ?, priority = ?, sprint_id = ?,
                assignee_id = ?, order_index = ?, start_date = ?, due_date = ?,
                estimated_hours = ?, actual_hours = ?, tags = ?, updated_at = ?
            WHERE id = ?
        """, (
            title, desc, status, priority, sprint_id,
            assignee_id, order_index, start_date, due_date,
            est_hours, act_hours, tags_json, now_str, task_id
        ))

        # 7. Subtasks if provided
        if isinstance(data.get("subtasks"), list):
            sub_row = conn.execute("SELECT COUNT(*) as cnt FROM subtasks WHERE task_id = ?", (task_id,)).fetchone()
            sub_count = sub_row["cnt"] if sub_row else 0
            for idx, sub_title in enumerate(data["subtasks"]):
                if isinstance(sub_title, str) and sub_title.strip():
                    cursor.execute("""
                        INSERT INTO subtasks (task_id, title, completed, order_index)
                        VALUES (?, ?, 0, ?)
                    """, (task_id, sub_title.strip(), sub_count + idx))

        # 8. Mapped Resources if provided
        if "resource_ids" in data and isinstance(data["resource_ids"], list):
            conn.execute("DELETE FROM task_resources WHERE task_id = ?", (task_id,))
            for r_id in data["resource_ids"]:
                try:
                    r_int = int(r_id)
                    cursor.execute("""
                        INSERT OR IGNORE INTO task_resources (task_id, resource_id, project_id, created_at)
                        VALUES (?, ?, ?, ?)
                    """, (task_id, r_int, project_id, now_str))
                except Exception:
                    pass

        changes = []
        if status != task["status"]:
            changes.append(f'status: {task["status"]} -> {status}')
        if priority != task["priority"]:
            changes.append(f'priority: {task["priority"]} -> {priority}')
        if start_date != task["start_date"]:
            changes.append(f'start_date: {task["start_date"]} -> {start_date}')
        if due_date != task["due_date"]:
            changes.append(f'due_date: {task["due_date"]} -> {due_date}')
        if assignee_id != task["assignee_id"]:
            changes.append(f'assignee_id: {task["assignee_id"]} -> {assignee_id}')
        
        details = ", ".join(changes) if changes else "Updated task fields"
        record_activity(conn, project_id, "User", "Task Updated", f'Task "{title}": {details}', task_id=task_id)

        # Trigger Action 4 (Task Assignment) if assignee changed
        if assignee_id and assignee_id != task["assignee_id"]:
            try:
                notify_task_assigned(conn, task_id, previous_assignee_id=task["assignee_id"])
            except Exception as e:
                print(f"[Notifier] Error sending assignment notification: {e}")

        # Trigger Action 5 (Task Completed) if status changed to done
        if status == "done" and task["status"] != "done":
            try:
                notify_task_completed(conn, task_id, actor_name="User")
            except Exception as e:
                print(f"[Notifier] Error sending completion notification: {e}")

        task_res = get_task_dict(conn, task_id)
        return json_response(task_res)

@app.post("/api/tasks/reorder")
def reorder_tasks():
    items = request.json or []
    if not isinstance(items, list):
        return json_response({"error": "Expected array of items"}, status=400)
    
    with get_db() as conn:
        for item in items:
            t_id = item.get("task_id")
            status = item.get("status")
            order_idx = item.get("new_order_index", 0)
            if t_id and status:
                prev_task = conn.execute("SELECT status FROM tasks WHERE id = ?", (t_id,)).fetchone()
                conn.execute("""
                    UPDATE tasks SET status = ?, order_index = ?, updated_at = ?
                    WHERE id = ?
                """, (status, order_idx, get_now_iso(), t_id))
                if status == "done" and prev_task and prev_task["status"] != "done":
                    try:
                        notify_task_completed(conn, t_id, actor_name="User")
                    except Exception as e:
                        print(f"[Notifier] Error sending completion notification: {e}")
        return json_response({"success": True})

@app.post("/api/tasks/<task_id:int>/move")
def move_task(task_id):
    data = request.json or {}
    direction = data.get("direction", "down")  # "up" or "down"
    
    with get_db() as conn:
        task = conn.execute("SELECT * FROM tasks WHERE id = ?", (task_id,)).fetchone()
        if not task:
            return json_response({"error": "Task not found"}, status=404)
        
        project_id = task["project_id"]
        tasks = conn.execute(
            "SELECT id, order_index FROM tasks WHERE project_id = ? ORDER BY order_index ASC, id ASC",
            (project_id,)
        ).fetchall()
        
        task_list = [dict(t) for t in tasks]
        idx = next((i for i, t in enumerate(task_list) if t["id"] == task_id), None)
        if idx is None:
            return json_response({"error": "Task not found in project"}, status=404)
        
        target_idx = idx - 1 if direction == "up" else idx + 1
        if 0 <= target_idx < len(task_list):
            task_list[idx], task_list[target_idx] = task_list[target_idx], task_list[idx]
            now_str = get_now_iso()
            for i, t in enumerate(task_list):
                conn.execute("UPDATE tasks SET order_index = ?, updated_at = ? WHERE id = ?", (i, now_str, t["id"]))
        
        bootstrap = get_bootstrap_payload(conn, active_project_id=project_id)
        return json_response({
            "success": True,
            "tasks": bootstrap["tasks"],
            "current_project": bootstrap["current_project"]
        })

@app.delete("/api/tasks/<task_id:int>")
def delete_task(task_id):
    with get_db() as conn:
        task = conn.execute("SELECT * FROM tasks WHERE id = ?", (task_id,)).fetchone()
        if not task:
            return json_response({"error": "Task not found"}, status=404)
        record_activity(conn, task["project_id"], "User", "Task Deleted", f'Deleted task "{task["title"]}"')
        conn.execute("DELETE FROM tasks WHERE id = ?", (task_id,))
        return json_response({"success": True})

# ==================== SUBTASKS ====================

@app.post("/api/tasks/<task_id:int>/subtasks")
def add_subtask(task_id):
    data = request.json or {}
    title = data.get("title", "").strip()
    if not title:
        return json_response({"error": "Subtask title is required"}, status=400)
    
    completed = 1 if data.get("completed") else 0
    with get_db() as conn:
        task = conn.execute("SELECT project_id FROM tasks WHERE id = ?", (task_id,)).fetchone()
        if not task:
            return json_response({"error": "Task not found"}, status=404)
        
        max_idx = conn.execute("SELECT COALESCE(MAX(order_index), -1) as m FROM subtasks WHERE task_id = ?", (task_id,)).fetchone()["m"]
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO subtasks (task_id, title, completed, order_index)
            VALUES (?, ?, ?, ?)
        """, (task_id, title, completed, max_idx + 1))
        s_id = cursor.lastrowid
        subtask = conn.execute("SELECT * FROM subtasks WHERE id = ?", (s_id,)).fetchone()
        return json_response(subtask, status=201)

@app.put("/api/subtasks/<subtask_id:int>")
def update_subtask(subtask_id):
    data = request.json or {}
    with get_db() as conn:
        sub = conn.execute("SELECT * FROM subtasks WHERE id = ?", (subtask_id,)).fetchone()
        if not sub:
            return json_response({"error": "Subtask not found"}, status=404)
        
        title = data.get("title", sub["title"])
        completed = (1 if data.get("completed") else 0) if "completed" in data else sub["completed"]

        conn.execute("UPDATE subtasks SET title = ?, completed = ? WHERE id = ?", (title, completed, subtask_id))
        updated = conn.execute("SELECT * FROM subtasks WHERE id = ?", (subtask_id,)).fetchone()
        return json_response(updated)

@app.delete("/api/subtasks/<subtask_id:int>")
def delete_subtask(subtask_id):
    with get_db() as conn:
        conn.execute("DELETE FROM subtasks WHERE id = ?", (subtask_id,))
        return json_response({"success": True})

# ==================== TIME LOGS ====================

@app.post("/api/tasks/<task_id:int>/timelogs")
def add_timelog(task_id):
    data = request.json or {}
    hours = float(data.get("hours") or 0.0)
    if hours <= 0:
        return json_response({"error": "Valid hours required"}, status=400)
    
    member_id = data.get("member_id")
    desc = data.get("description", "")
    logged_date = data.get("logged_date") or datetime.now(timezone.utc).strftime("%Y-%m-%d")
    now_str = get_now_iso()

    with get_db() as conn:
        task = conn.execute("SELECT project_id, title, assignee_id FROM tasks WHERE id = ?", (task_id,)).fetchone()
        if not task:
            return json_response({"error": "Task not found"}, status=404)
        
        if not member_id:
            if task["assignee_id"]:
                member_id = task["assignee_id"]
            else:
                first_m = conn.execute("SELECT id FROM members WHERE project_id = ? LIMIT 1", (task["project_id"],)).fetchone()
                if first_m:
                    member_id = first_m["id"]

        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO timelogs (task_id, member_id, hours, description, logged_date, created_at)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (task_id, member_id, hours, desc, logged_date, now_str))
        log_id = cursor.lastrowid

        total_hours = conn.execute("SELECT COALESCE(SUM(hours), 0) as s FROM timelogs WHERE task_id = ?", (task_id,)).fetchone()["s"]
        conn.execute("UPDATE tasks SET actual_hours = ?, updated_at = ? WHERE id = ?", (total_hours, now_str, task_id))

        record_activity(conn, task["project_id"], "User", "Time Logged", f'Logged {hours}h on "{task["title"]}"', task_id=task_id)

        log_row = conn.execute("""
            SELECT tl.*, m.name as member_name, m.avatar_color as member_avatar
            FROM timelogs tl
            LEFT JOIN members m ON tl.member_id = m.id
            WHERE tl.id = ?
        """, (log_id,)).fetchone()
        return json_response(log_row, status=201)

@app.delete("/api/timelogs/<log_id:int>")
def delete_timelog(log_id):
    with get_db() as conn:
        log = conn.execute("SELECT * FROM timelogs WHERE id = ?", (log_id,)).fetchone()
        if not log:
            return json_response({"error": "Time log not found"}, status=404)
        
        task_id = log["task_id"]
        conn.execute("DELETE FROM timelogs WHERE id = ?", (log_id,))
        
        total_hours = conn.execute("SELECT COALESCE(SUM(hours), 0) as s FROM timelogs WHERE task_id = ?", (task_id,)).fetchone()["s"]
        conn.execute("UPDATE tasks SET actual_hours = ?, updated_at = ? WHERE id = ?", (total_hours, get_now_iso(), task_id))
        return json_response({"success": True})

@app.get("/api/projects/<project_id:int>/timelogs")
def get_project_timelogs(project_id):
    with get_db() as conn:
        logs = conn.execute("""
            SELECT tl.*, t.title as task_title, t.priority as task_priority, m.name as member_name, m.avatar_color as member_avatar
            FROM timelogs tl
            JOIN tasks t ON tl.task_id = t.id
            LEFT JOIN members m ON tl.member_id = m.id
            WHERE t.project_id = ?
            ORDER BY tl.logged_date DESC, tl.id DESC
        """, (project_id,)).fetchall()
        return json_response(logs)

# ==================== ANALYTICS ====================

@app.get("/api/portfolio/analytics")
@app.get("/api/analytics/portfolio")
def get_portfolio_analytics():
    with get_db() as conn:
        projects = conn.execute("""
            SELECT
                p.id, p.name, p.description, p.color, p.created_at, p.updated_at,
                COUNT(t.id) as total_tasks,
                SUM(CASE WHEN t.status = 'done' THEN 1 ELSE 0 END) as done_tasks,
                SUM(CASE WHEN t.status = 'in_progress' THEN 1 ELSE 0 END) as in_progress_tasks,
                SUM(CASE WHEN t.status = 'todo' THEN 1 ELSE 0 END) as todo_tasks,
                SUM(CASE WHEN t.status != 'done' AND t.due_date < date('now') AND t.due_date IS NOT NULL AND t.due_date != '' THEN 1 ELSE 0 END) as overdue_tasks,
                COALESCE(SUM(t.estimated_hours), 0) as total_est_hours,
                COALESCE(SUM(t.actual_hours), 0) as total_act_hours
            FROM projects p
            LEFT JOIN tasks t ON t.project_id = p.id
            GROUP BY p.id
            ORDER BY p.updated_at DESC
        """).fetchall()

        projects_list = []
        tot_tasks = 0
        tot_done = 0
        tot_inprogress = 0
        tot_overdue = 0
        tot_est_hours = 0.0
        tot_act_hours = 0.0

        for p in projects:
            p_dict = dict(p)
            p_tot = p_dict["total_tasks"] or 0
            p_done = p_dict["done_tasks"] or 0
            p_overdue = p_dict["overdue_tasks"] or 0
            p_rate = round((p_done / p_tot * 100), 1) if p_tot > 0 else 0.0
            p_dict["completion_rate"] = p_rate
            
            # Delivery health classification
            if p_overdue > 0:
                p_dict["health_status"] = "overdue"
            elif p_rate < 30 and p_tot > 5:
                p_dict["health_status"] = "at_risk"
            else:
                p_dict["health_status"] = "on_track"

            projects_list.append(p_dict)

            tot_tasks += p_tot
            tot_done += p_done
            tot_inprogress += (p_dict["in_progress_tasks"] or 0)
            tot_overdue += p_overdue
            tot_est_hours += float(p_dict["total_est_hours"] or 0.0)
            tot_act_hours += float(p_dict["total_act_hours"] or 0.0)

        overall_completion_rate = round((tot_done / tot_tasks * 100), 1) if tot_tasks > 0 else 0.0

        # Global status distribution
        status_counts = conn.execute("""
            SELECT status, COUNT(*) as count, COALESCE(SUM(estimated_hours), 0) as est_hours, COALESCE(SUM(actual_hours), 0) as act_hours
            FROM tasks
            GROUP BY status
        """).fetchall()

        # Global priority distribution
        priority_counts = conn.execute("""
            SELECT priority, COUNT(*) as count
            FROM tasks
            GROUP BY priority
        """).fetchall()

        # Cross-project team workload
        workload_rows = conn.execute("""
            SELECT 
                COALESCE(NULLIF(TRIM(m.name), ''), 'Unassigned') as name,
                COALESCE(NULLIF(TRIM(m.role), ''), 'Team Member') as role,
                COALESCE(m.avatar_color, '#3B82F6') as avatar_color,
                COUNT(t.id) as assigned_tasks,
                COUNT(CASE WHEN t.status = 'done' THEN 1 END) as completed_tasks,
                COALESCE(SUM(t.estimated_hours), 0) as total_est_hours,
                COALESCE(SUM(t.actual_hours), 0) as total_act_hours
            FROM tasks t
            LEFT JOIN members m ON t.assignee_id = m.id
            GROUP BY LOWER(TRIM(COALESCE(m.name, 'Unassigned')))
            ORDER BY assigned_tasks DESC
        """).fetchall()
        workload = [dict(r) for r in workload_rows]

        # Recent activities across all projects
        acts = conn.execute("""
            SELECT a.*, p.name as project_name, t.title as task_title
            FROM activity_logs a
            LEFT JOIN projects p ON a.project_id = p.id
            LEFT JOIN tasks t ON a.task_id = t.id
            ORDER BY a.timestamp DESC
            LIMIT 25
        """).fetchall()

        result = {
            "kpis": {
                "total_projects": len(projects_list),
                "total_tasks": tot_tasks,
                "done_tasks": tot_done,
                "in_progress_tasks": tot_inprogress,
                "overdue_tasks": tot_overdue,
                "completion_rate": overall_completion_rate,
                "total_est_hours": tot_est_hours,
                "total_act_hours": tot_act_hours
            },
            "projects": projects_list,
            "status_distribution": [dict(r) for r in status_counts],
            "priority_distribution": [dict(r) for r in priority_counts],
            "workload": workload,
            "activities": [dict(r) for r in acts]
        }
        return json_response(result)

@app.get("/api/projects/<project_id:int>/analytics")
def get_analytics(project_id):
    sprint_id = request.query.get("sprint_id")
    with get_db() as conn:
        project = conn.execute("SELECT * FROM projects WHERE id = ?", (project_id,)).fetchone()
        if not project:
            return json_response({"error": "Project not found"}, status=404)

        if sprint_id and sprint_id.isdigit():
            sprint = conn.execute("SELECT * FROM sprints WHERE id = ? AND project_id = ?", (int(sprint_id), project_id)).fetchone()
        else:
            sprint = conn.execute("SELECT * FROM sprints WHERE project_id = ? AND status = 'active' ORDER BY start_date DESC LIMIT 1", (project_id,)).fetchone()
            if not sprint:
                sprint = conn.execute("SELECT * FROM sprints WHERE project_id = ? ORDER BY start_date DESC LIMIT 1", (project_id,)).fetchone()

        status_counts = conn.execute("""
            SELECT status, COUNT(*) as count, COALESCE(SUM(estimated_hours), 0) as est_hours, COALESCE(SUM(actual_hours), 0) as act_hours
            FROM tasks
            WHERE project_id = ?
            GROUP BY status
        """, (project_id,)).fetchall()

        priority_counts = conn.execute("""
            SELECT priority, COUNT(*) as count
            FROM tasks
            WHERE project_id = ?
            GROUP BY priority
        """, (project_id,)).fetchall()

        workload = conn.execute("""
            SELECT m.id, m.name, m.role, m.avatar_color,
                COUNT(t.id) as assigned_tasks,
                COUNT(CASE WHEN t.status = 'done' THEN 1 END) as completed_tasks,
                COALESCE(SUM(t.estimated_hours), 0) as total_est_hours,
                COALESCE(SUM(t.actual_hours), 0) as total_act_hours
            FROM members m
            LEFT JOIN tasks t ON m.id = t.assignee_id AND t.project_id = ?
            WHERE m.project_id = ?
            GROUP BY m.id
            ORDER BY assigned_tasks DESC
        """, (project_id, project_id)).fetchall()

        if not workload:
            workload_rows = conn.execute("""
                SELECT 
                    COALESCE(m.id, 0) as id,
                    COALESCE(m.name, 'Unassigned') as name,
                    COALESCE(m.role, 'Team Member') as role,
                    COALESCE(m.avatar_color, '#3B82F6') as avatar_color,
                    COUNT(t.id) as assigned_tasks,
                    COUNT(CASE WHEN t.status = 'done' THEN 1 END) as completed_tasks,
                    COALESCE(SUM(t.estimated_hours), 0) as total_est_hours,
                    COALESCE(SUM(t.actual_hours), 0) as total_act_hours
                FROM tasks t
                LEFT JOIN members m ON t.assignee_id = m.id
                WHERE t.project_id = ?
                GROUP BY COALESCE(m.id, 0)
                ORDER BY assigned_tasks DESC
            """, (project_id,)).fetchall()
            workload = [dict(r) for r in workload_rows]
        else:
            workload = [dict(r) for r in workload]

        kpi_row = conn.execute("""
            SELECT
                COUNT(*) as total_tasks,
                COUNT(CASE WHEN status = 'done' THEN 1 END) as done_tasks,
                COUNT(CASE WHEN status = 'in_progress' THEN 1 END) as in_progress_tasks,
                COUNT(CASE WHEN status != 'done' AND due_date < date('now') AND due_date IS NOT NULL THEN 1 END) as overdue_tasks,
                COALESCE(SUM(estimated_hours), 0) as total_est_hours,
                COALESCE(SUM(actual_hours), 0) as total_act_hours
            FROM tasks
            WHERE project_id = ?
        """, (project_id,)).fetchone()

        total = kpi_row["total_tasks"] or 0
        done = kpi_row["done_tasks"] or 0
        completion_rate = round((done / total * 100), 1) if total > 0 else 0

        burndown_data = {"labels": [], "ideal": [], "actual": [], "sprint_name": sprint["name"] if sprint else "Project Timeline"}

        if sprint and sprint["start_date"] and sprint["end_date"]:
            try:
                start = datetime.strptime(sprint["start_date"][:10], "%Y-%m-%d")
                end = datetime.strptime(sprint["end_date"][:10], "%Y-%m-%d")
                days_total = max((end - start).days, 1)
                
                sprint_tasks_total_est = conn.execute(
                    "SELECT COALESCE(SUM(estimated_hours), 0) as s FROM tasks WHERE sprint_id = ?",
                    (sprint["id"],)
                ).fetchone()["s"] or 40.0

                labels = []
                ideal = []
                actual = []
                today = datetime.now()

                for i in range(days_total + 1):
                    day_dt = start + timedelta(days=i)
                    labels.append(day_dt.strftime("%b %d"))
                    
                    ideal_val = round(sprint_tasks_total_est * (1 - (i / days_total)), 1)
                    ideal.append(max(ideal_val, 0))

                    if day_dt.date() <= (today + timedelta(days=1)).date():
                        completed_by_day = conn.execute("""
                            SELECT COALESCE(SUM(hours), 0) as s FROM timelogs tl
                            JOIN tasks t ON tl.task_id = t.id
                            WHERE t.sprint_id = ? AND tl.logged_date <= ?
                        """, (sprint["id"], day_dt.strftime("%Y-%m-%d"))).fetchone()["s"]
                        actual_val = max(round(sprint_tasks_total_est - completed_by_day, 1), 0)
                        actual.append(actual_val)

                burndown_data = {
                    "labels": labels,
                    "ideal": ideal,
                    "actual": actual,
                    "sprint_name": sprint["name"]
                }
            except Exception:
                pass

        if not burndown_data["labels"]:
            dates_row = conn.execute("""
                SELECT MIN(start_date) as min_start, MAX(due_date) as max_due,
                       COALESCE(SUM(estimated_hours), 0) as total_est,
                       COALESCE(SUM(actual_hours), 0) as total_act
                FROM tasks
                WHERE project_id = ? AND start_date IS NOT NULL
            """, (project_id,)).fetchone()
            
            if dates_row and dates_row["min_start"] and dates_row["max_due"]:
                try:
                    start = datetime.strptime(dates_row["min_start"][:10], "%Y-%m-%d")
                    end = datetime.strptime(dates_row["max_due"][:10], "%Y-%m-%d")
                    total_est = float(dates_row["total_est"] or 40.0)
                    days_total = max((end - start).days, 1)
                    step_days = max(days_total // 7, 1)
                    
                    labels = []
                    ideal = []
                    actual = []
                    today = datetime.now()
                    num_steps = max(days_total // step_days, 1)
                    
                    for i in range(num_steps + 1):
                        day_dt = start + timedelta(days=min(i * step_days, days_total))
                        labels.append(day_dt.strftime("%b %d"))
                        ideal_val = round(total_est * (1 - (min(i * step_days, days_total) / days_total)), 1)
                        ideal.append(max(ideal_val, 0))
                        
                        if day_dt.date() <= (today + timedelta(days=1)).date():
                            completed_sum = conn.execute("""
                                SELECT COALESCE(SUM(estimated_hours), 0) as s FROM tasks
                                WHERE project_id = ? AND status = 'done' AND (due_date <= ? OR start_date <= ?)
                            """, (project_id, day_dt.strftime("%Y-%m-%d"), day_dt.strftime("%Y-%m-%d"))).fetchone()["s"]
                            actual.append(max(round(total_est - completed_sum, 1), 0))
                    
                    burndown_data = {
                        "labels": labels,
                        "ideal": ideal,
                        "actual": actual,
                        "sprint_name": "Project Lifecycle Progress"
                    }
                except Exception:
                    pass

        dept_counts = conn.execute("""
            SELECT 
                COALESCE(NULLIF(TRIM(m.role), ''), 'Unassigned') as department,
                COUNT(t.id) as count,
                COALESCE(SUM(t.estimated_hours), 0) as est_hours,
                COUNT(CASE WHEN t.status = 'done' THEN 1 END) as completed_count
            FROM tasks t
            LEFT JOIN members m ON t.assignee_id = m.id
            WHERE t.project_id = ?
            GROUP BY department
            ORDER BY count DESC
        """, (project_id,)).fetchall()

        result = {
            "kpis": {
                **dict(kpi_row),
                "completion_rate": completion_rate,
                "sprint_name": sprint["name"] if sprint else "Project Lifecycle"
            },
            "status_distribution": [dict(r) for r in status_counts],
            "priority_distribution": [dict(r) for r in priority_counts],
            "department_distribution": [dict(r) for r in dept_counts],
            "workload": workload,
            "burndown": burndown_data
        }
        return json_response(result)

@app.get("/api/projects/<project_id:int>/activity")
def get_activity(project_id):
    with get_db() as conn:
        acts = conn.execute("""
            SELECT a.*, t.title as task_title
            FROM activity_logs a
            LEFT JOIN tasks t ON a.task_id = t.id
            WHERE a.project_id = ?
            ORDER BY a.timestamp DESC
            LIMIT 40
        """, (project_id,)).fetchall()
        return json_response(acts)

# ==================== CUMULATIVE PROJECT ACTIVITY REPORT ====================

def build_cumulative_project_report(conn, project_id, current_user=None):
    project = conn.execute("SELECT * FROM projects WHERE id = ?", (project_id,)).fetchone()
    if not project:
        return None

    # Fetch all activities with assignee, sprint, subtasks, timelogs
    raw_tasks = conn.execute("""
        SELECT t.*,
               m.name as assignee_name, m.role as assignee_role, m.avatar_color as assignee_avatar, m.email as assignee_email,
               s.name as sprint_name
        FROM tasks t
        LEFT JOIN members m ON t.assignee_id = m.id
        LEFT JOIN sprints s ON t.sprint_id = s.id
        WHERE t.project_id = ?
        ORDER BY t.order_index ASC, t.id ASC
    """, (project_id,)).fetchall()

    members = conn.execute("SELECT * FROM members WHERE project_id = ? ORDER BY name ASC", (project_id,)).fetchall()
    milestones = conn.execute("SELECT * FROM milestones WHERE project_id = ? ORDER BY due_date ASC", (project_id,)).fetchall()
    sprints = conn.execute("SELECT * FROM sprints WHERE project_id = ? ORDER BY start_date ASC", (project_id,)).fetchall()

    # Subtasks
    subtasks_raw = conn.execute("""
        SELECT st.*
        FROM subtasks st
        JOIN tasks t ON st.task_id = t.id
        WHERE t.project_id = ?
        ORDER BY st.task_id, st.order_index ASC, st.id ASC
    """, (project_id,)).fetchall()
    subtasks_by_task = {}
    for st in subtasks_raw:
        tid = st["task_id"]
        subtasks_by_task.setdefault(tid, []).append({
            "id": st["id"],
            "title": st["title"],
            "completed": bool(st["completed"]),
            "order_index": st["order_index"]
        })

    # Timelogs
    timelogs_raw = conn.execute("""
        SELECT tl.*, m.name as member_name
        FROM timelogs tl
        JOIN tasks t ON tl.task_id = t.id
        LEFT JOIN members m ON tl.member_id = m.id
        WHERE t.project_id = ?
        ORDER BY tl.logged_date DESC
    """, (project_id,)).fetchall()
    timelogs_by_task = {}
    for tl in timelogs_raw:
        tid = tl["task_id"]
        timelogs_by_task.setdefault(tid, []).append({
            "id": tl["id"],
            "member_id": tl["member_id"],
            "member_name": tl["member_name"] or "Team Member",
            "hours": safe_float(tl["hours"]),
            "description": tl["description"] or "",
            "logged_date": tl["logged_date"]
        })

    # Activity Logs
    activity_logs_raw = conn.execute("""
        SELECT * FROM activity_logs
        WHERE project_id = ?
        ORDER BY timestamp DESC
    """, (project_id,)).fetchall()
    activity_by_task = {}
    for al in activity_logs_raw:
        tid = al["task_id"]
        if tid:
            activity_by_task.setdefault(tid, []).append({
                "id": al["id"],
                "user_name": al["user_name"],
                "action": al["action"],
                "details": al["details"] or "",
                "timestamp": al["timestamp"]
            })

    today_str = datetime.now(timezone.utc).strftime("%Y-%m-%d")

    total_activities = len(raw_tasks)
    completed_count = 0
    in_progress_count = 0
    todo_count = 0
    in_review_count = 0
    backlog_count = 0
    overdue_count = 0
    urgent_high_count = 0
    total_est_hours = 0.0
    total_act_hours = 0.0
    with_dates_count = 0
    unassigned_count = 0

    status_counts = {"done": 0, "in_progress": 0, "todo": 0, "in_review": 0, "backlog": 0}
    priority_counts = {"urgent": 0, "high": 0, "medium": 0, "low": 0}

    assignee_stats = {}
    for m in members:
        assignee_stats[m["id"]] = {
            "id": m["id"],
            "name": m["name"],
            "role": m["role"] or "Member",
            "email": m["email"] or "",
            "avatar_color": m["avatar_color"] or "#6366F1",
            "total": 0,
            "completed": 0,
            "in_progress": 0,
            "pending": 0,
            "overdue": 0,
            "total_est_hours": 0.0,
            "total_act_hours": 0.0
        }
    unassigned_stats = {
        "id": None,
        "name": "Unassigned",
        "role": "Not Assigned",
        "email": "",
        "avatar_color": "#94A3B8",
        "total": 0,
        "completed": 0,
        "in_progress": 0,
        "pending": 0,
        "overdue": 0,
        "total_est_hours": 0.0,
        "total_act_hours": 0.0
    }

    overdue_list = []
    urgent_high_pending_list = []
    missing_dates_list = []
    unassigned_list = []
    tasks_output = []

    for idx, t in enumerate(raw_tasks):
        tid = t["id"]
        status = t["status"] or "todo"
        priority = t["priority"] or "medium"
        est_h = safe_float(t["estimated_hours"])
        act_h = safe_float(t["actual_hours"])
        total_est_hours += est_h
        total_act_hours += act_h

        tags_list = []
        if t["tags"]:
            try:
                tags_list = json.loads(t["tags"]) if isinstance(t["tags"], str) else t["tags"]
                if not isinstance(tags_list, list):
                    tags_list = [str(tags_list)]
            except Exception:
                tags_list = [tag.strip() for tag in str(t["tags"]).split(",") if tag.strip()]

        subtasks = subtasks_by_task.get(tid, [])
        st_completed = sum(1 for s in subtasks if s["completed"])
        st_count = len(subtasks)

        if status == "done":
            progress_pct = 100
        elif status == "in_review":
            progress_pct = 85
        elif status == "in_progress":
            progress_pct = max(10, int((st_completed / st_count) * 100)) if st_count > 0 else 50
        else:
            progress_pct = 0

        start_d = t["start_date"]
        due_d = t["due_date"]
        duration_days = None
        is_overdue = False
        delay_days = 0

        if start_d and due_d:
            try:
                d1 = datetime.strptime(start_d, "%Y-%m-%d")
                d2 = datetime.strptime(due_d, "%Y-%m-%d")
                duration_days = max(1, (d2 - d1).days + 1)
            except Exception:
                pass

        if due_d:
            with_dates_count += 1
            if status != "done" and due_d < today_str:
                is_overdue = True
                overdue_count += 1
                try:
                    due_obj = datetime.strptime(due_d, "%Y-%m-%d")
                    now_obj = datetime.strptime(today_str, "%Y-%m-%d")
                    delay_days = max(1, (now_obj - due_obj).days)
                except Exception:
                    delay_days = 1
        else:
            if not start_d and not due_d:
                missing_dates_list.append({
                    "id": tid,
                    "seq_num": idx + 1,
                    "title": t["title"],
                    "assignee_name": t["assignee_name"] or "Unassigned",
                    "status": status,
                    "priority": priority
                })

        if status == "done":
            completed_count += 1
            status_counts["done"] = status_counts.get("done", 0) + 1
        elif status == "in_progress":
            in_progress_count += 1
            status_counts["in_progress"] = status_counts.get("in_progress", 0) + 1
        elif status == "in_review":
            in_review_count += 1
            status_counts["in_review"] = status_counts.get("in_review", 0) + 1
        elif status == "backlog":
            backlog_count += 1
            status_counts["backlog"] = status_counts.get("backlog", 0) + 1
        else:
            todo_count += 1
            status_counts["todo"] = status_counts.get("todo", 0) + 1

        priority_counts[priority] = priority_counts.get(priority, 0) + 1
        if priority in ("urgent", "high"):
            urgent_high_count += 1
            if status != "done":
                urgent_high_pending_list.append({
                    "id": tid,
                    "seq_num": idx + 1,
                    "title": t["title"],
                    "priority": priority,
                    "status": status,
                    "assignee_name": t["assignee_name"] or "Unassigned",
                    "due_date": due_d or "Not Declared"
                })

        if is_overdue:
            overdue_list.append({
                "id": tid,
                "seq_num": idx + 1,
                "title": t["title"],
                "due_date": due_d,
                "delay_days": delay_days,
                "assignee_name": t["assignee_name"] or "Unassigned",
                "priority": priority,
                "status": status
            })

        aid = t["assignee_id"]
        if aid and aid in assignee_stats:
            astat = assignee_stats[aid]
            astat["total"] += 1
            astat["total_est_hours"] += est_h
            astat["total_act_hours"] += act_h
            if status == "done":
                astat["completed"] += 1
            elif status == "in_progress":
                astat["in_progress"] += 1
            else:
                astat["pending"] += 1
            if is_overdue:
                astat["overdue"] += 1
        else:
            unassigned_count += 1
            unassigned_stats["total"] += 1
            unassigned_stats["total_est_hours"] += est_h
            unassigned_stats["total_act_hours"] += act_h
            if status == "done":
                unassigned_stats["completed"] += 1
            elif status == "in_progress":
                unassigned_stats["in_progress"] += 1
            else:
                unassigned_stats["pending"] += 1
            if is_overdue:
                unassigned_stats["overdue"] += 1
            unassigned_list.append({
                "id": tid,
                "seq_num": idx + 1,
                "title": t["title"],
                "status": status,
                "priority": priority,
                "due_date": due_d or "Not Declared"
            })

        deliverables_text = None
        if t["description"] and "deliverable" in t["description"].lower():
            deliverables_text = t["description"]
        elif tags_list:
            deliverables_text = f"Activity Deliverables Tagged: {', '.join(tags_list)}"

        depends_on = []
        blocks = []
        if idx > 0:
            prev_t = raw_tasks[idx - 1]
            depends_on.append({
                "id": prev_t["id"],
                "seq_num": idx,
                "title": prev_t["title"],
                "status": prev_t["status"]
            })
        if idx < len(raw_tasks) - 1:
            next_t = raw_tasks[idx + 1]
            blocks.append({
                "id": next_t["id"],
                "seq_num": idx + 2,
                "title": next_t["title"],
                "status": next_t["status"]
            })

        tasks_output.append({
            "id": tid,
            "seq_num": idx + 1,
            "title": t["title"],
            "description": t["description"] or "",
            "status": status,
            "priority": priority,
            "order_index": t["order_index"],
            "start_date": start_d,
            "due_date": due_d,
            "duration_days": duration_days,
            "is_overdue": is_overdue,
            "delay_days": delay_days,
            "estimated_hours": est_h,
            "actual_hours": act_h,
            "progress_pct": progress_pct,
            "assignee_id": aid,
            "assignee_name": t["assignee_name"] or "Unassigned",
            "assignee_role": t["assignee_role"] or "Not Set",
            "assignee_avatar": t["assignee_avatar"] or "#94A3B8",
            "assignee_email": t["assignee_email"] or "",
            "sprint_id": t["sprint_id"],
            "sprint_name": t["sprint_name"] or "Standard Phase",
            "tags": tags_list,
            "subtasks": subtasks,
            "subtask_count": st_count,
            "subtask_completed_count": st_completed,
            "timelogs": timelogs_by_task.get(tid, []),
            "activity_history": activity_by_task.get(tid, []),
            "deliverables": deliverables_text or "",
            "dependencies": {
                "depends_on": depends_on,
                "blocks": blocks
            },
            "created_at": t["created_at"],
            "updated_at": t["updated_at"]
        })

    completion_pct = round((completed_count / total_activities * 100), 1) if total_activities > 0 else 0.0

    assignee_summary_list = []
    for uid, s in assignee_stats.items():
        if s["total"] > 0:
            s["completion_pct"] = round((s["completed"] / s["total"] * 100), 1)
            assignee_summary_list.append(s)
    if unassigned_stats["total"] > 0:
        unassigned_stats["completion_pct"] = round((unassigned_stats["completed"] / unassigned_stats["total"] * 100), 1)
        assignee_summary_list.append(unassigned_stats)

    assignee_summary_list.sort(key=lambda x: x["total"], reverse=True)

    gen_by = "System Administrator"
    if current_user and isinstance(current_user, dict) and current_user.get("full_name"):
        gen_by = current_user["full_name"]

    first_start = next((t["start_date"] for t in tasks_output if t["start_date"]), "Project Inception")
    last_due = next((t["due_date"] for t in reversed(tasks_output) if t["due_date"]), "Current Date")

    now_iso = get_now_iso()
    period_str = f"{first_start} to {last_due}"

    return {
        "project": {
            "id": project["id"],
            "name": project["name"],
            "description": project["description"] or "",
            "created_at": project["created_at"],
            "updated_at": project["updated_at"]
        },
        "metadata": {
            "generated_at": now_iso,
            "generated_by": gen_by,
            "reporting_period": period_str,
            "project_name": project["name"],
            "project_id": project["id"]
        },
        "generated_at": now_iso,
        "generated_by": gen_by,
        "reporting_period": period_str,
        "kpis": {
            "total_activities": total_activities,
            "completed": completed_count,
            "in_progress": in_progress_count,
            "to_do": todo_count,
            "in_review": in_review_count,
            "backlog": backlog_count,
            "overdue": overdue_count,
            "urgent_high": urgent_high_count,
            "completion_pct": completion_pct,
            "total_estimated_hours": total_est_hours,
            "total_actual_hours": total_act_hours,
            "activities_with_dates": with_dates_count,
            "activities_without_dates": total_activities - with_dates_count,
            "unassigned_activities": unassigned_count
        },
        "status_breakdown": status_counts,
        "priority_breakdown": priority_counts,
        "assignee_summary": assignee_summary_list,
        "workload": assignee_summary_list,
        "milestones": [dict(m) for m in milestones],
        "sprints": [dict(s) for s in sprints],
        "attention_required": {
            "overdue": overdue_list,
            "urgent_high_pending": urgent_high_pending_list,
            "missing_dates": missing_dates_list,
            "unassigned": unassigned_list
        },
        "activities": tasks_output
    }

def generate_cumulative_project_excel(report):
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()

    navy_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    blue_header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    accent_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    light_blue_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    status_done_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    status_prog_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    prio_urgent_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    prio_high_fill = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")

    title_font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    section_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True, color="0F172A")
    regular_font = Font(name="Calibri", size=10, color="334155")
    small_italic_font = Font(name="Calibri", size=9, italic=True, color="64748B")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # 1. Executive Summary Sheet
    ws1 = wb.active
    ws1.title = "Executive Summary"
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells("A1:G2")
    ws1["A1"] = f"ProjectPulse — Cumulative Project Activity Report\n{report['project']['name']} (ID: #{report['project']['id']})"
    ws1["A1"].font = title_font
    ws1["A1"].fill = navy_fill
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws1["A3"] = f"Generated: {report['generated_at'][:19].replace('T', ' ')} UTC | Generated By: {report['generated_by']} | Period: {report.get('reporting_period', 'All')}"
    ws1["A3"].font = small_italic_font
    ws1.merge_cells("A3:G3")

    ws1["A5"] = "1. Executive Key Performance Indicators (KPIs)"
    ws1["A5"].font = section_font
    ws1["A5"].fill = accent_fill
    ws1.merge_cells("A5:G5")

    kpi_items = [
        ("Total Activities", report["kpis"]["total_activities"], "Completed Activities", report["kpis"]["completed"]),
        ("In Progress Activities", report["kpis"]["in_progress"], "Pending / To Do", report["kpis"]["to_do"]),
        ("Overdue Activities", report["kpis"]["overdue"], "Urgent & High Priority", report["kpis"]["urgent_high"]),
        ("Overall Completion Rate", f"{report['kpis']['completion_pct']}%", "Activities with Dates", report["kpis"]["activities_with_dates"]),
        ("Total Estimated Hours", f"{report['kpis']['total_estimated_hours']} hrs", "Total Actual Hours Logged", f"{report['kpis']['total_actual_hours']} hrs"),
    ]

    r_idx = 6
    for left_label, left_val, right_label, right_val in kpi_items:
        ws1.cell(row=r_idx, column=1, value=left_label).font = bold_font
        ws1.cell(row=r_idx, column=2, value=left_val).font = bold_font
        ws1.cell(row=r_idx, column=4, value=right_label).font = bold_font
        ws1.cell(row=r_idx, column=5, value=right_val).font = bold_font
        for c in [1, 2, 4, 5]:
            ws1.cell(row=r_idx, column=c).border = thin_border
            if c in [1, 4]:
                ws1.cell(row=r_idx, column=c).fill = light_blue_fill
        r_idx += 1

    r_idx += 2
    ws1.cell(row=r_idx, column=1, value="2. Team Resource Workload & Delivery Summary").font = section_font
    ws1.cell(row=r_idx, column=1).fill = accent_fill
    ws1.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=7)
    r_idx += 1

    team_headers = ["Assignee Name", "Role", "Assigned Activities", "Completed", "In Progress", "Overdue", "Completion %"]
    for col_idx, h in enumerate(team_headers, 1):
        cell = ws1.cell(row=r_idx, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = blue_header_fill
        cell.alignment = Alignment(horizontal="center" if col_idx > 2 else "left")
        cell.border = thin_border
    r_idx += 1

    for row_data in report.get("assignee_summary", []):
        ws1.cell(row=r_idx, column=1, value=row_data["name"]).font = bold_font
        ws1.cell(row=r_idx, column=2, value=row_data["role"]).font = regular_font
        ws1.cell(row=r_idx, column=3, value=row_data["total"]).alignment = Alignment(horizontal="center")
        ws1.cell(row=r_idx, column=4, value=row_data["completed"]).alignment = Alignment(horizontal="center")
        ws1.cell(row=r_idx, column=5, value=row_data["in_progress"]).alignment = Alignment(horizontal="center")
        ws1.cell(row=r_idx, column=6, value=row_data["overdue"]).alignment = Alignment(horizontal="center")
        ws1.cell(row=r_idx, column=7, value=f"{row_data.get('completion_pct', 0)}%").alignment = Alignment(horizontal="center")
        for c in range(1, 8):
            ws1.cell(row=r_idx, column=c).border = thin_border
        r_idx += 1

    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # 2. Activity Register Sheet
    ws2 = wb.create_sheet(title="Activity Register")
    ws2.views.sheetView[0].showGridLines = True

    act_headers = [
        "Seq #", "ID", "Activity Title", "Status", "Priority",
        "Owner", "Role", "Start Date", "End Date", "Duration (Days)",
        "Progress %", "Est Hours", "Actual Hours", "Subtasks (Done/Total)",
        "Tags", "Description", "Deliverables", "Dependencies (Depends On)", "Dependencies (Blocks)"
    ]

    for col_idx, h in enumerate(act_headers, 1):
        cell = ws2.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = blue_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws2.row_dimensions[1].height = 28

    r_idx = 2
    for t in report.get("activities", []):
        subtasks_str = f"{t.get('subtask_completed_count', 0)} / {t.get('subtask_count', 0)}" if t.get('subtask_count', 0) > 0 else "None"
        tags_str = ", ".join(t.get("tags", [])) if t.get("tags") else "None"
        depends_on_str = ", ".join([f"#{d['seq_num']} {d['title']}" for d in t.get("dependencies", {}).get("depends_on", [])]) or "None"
        blocks_str = ", ".join([f"#{b['seq_num']} {b['title']}" for b in t.get("dependencies", {}).get("blocks", [])]) or "None"

        ws2.cell(row=r_idx, column=1, value=f"#{t['seq_num']:02d}").alignment = Alignment(horizontal="center")
        ws2.cell(row=r_idx, column=2, value=t["id"]).alignment = Alignment(horizontal="center")
        ws2.cell(row=r_idx, column=3, value=t["title"]).font = bold_font
        
        status_cell = ws2.cell(row=r_idx, column=4, value=t["status"].upper().replace("_", " "))
        status_cell.alignment = Alignment(horizontal="center")
        if t["status"] == "done":
            status_cell.fill = status_done_fill
        elif t["status"] == "in_progress":
            status_cell.fill = status_prog_fill
        
        prio_cell = ws2.cell(row=r_idx, column=5, value=t["priority"].upper())
        prio_cell.alignment = Alignment(horizontal="center")
        if t["priority"] == "urgent":
            prio_cell.fill = prio_urgent_fill
        elif t["priority"] == "high":
            prio_cell.fill = prio_high_fill

        ws2.cell(row=r_idx, column=6, value=t["assignee_name"])
        ws2.cell(row=r_idx, column=7, value=t["assignee_role"])
        ws2.cell(row=r_idx, column=8, value=t["start_date"] or "Not Set").alignment = Alignment(horizontal="center")
        ws2.cell(row=r_idx, column=9, value=t["due_date"] or "Not Set").alignment = Alignment(horizontal="center")
        ws2.cell(row=r_idx, column=10, value=t["duration_days"] or "N/A").alignment = Alignment(horizontal="center")
        ws2.cell(row=r_idx, column=11, value=f"{t['progress_pct']}%").alignment = Alignment(horizontal="center")
        ws2.cell(row=r_idx, column=12, value=t["estimated_hours"]).alignment = Alignment(horizontal="center")
        ws2.cell(row=r_idx, column=13, value=t["actual_hours"]).alignment = Alignment(horizontal="center")
        ws2.cell(row=r_idx, column=14, value=subtasks_str).alignment = Alignment(horizontal="center")
        ws2.cell(row=r_idx, column=15, value=tags_str)
        ws2.cell(row=r_idx, column=16, value=t["description"] or "Not Available")
        ws2.cell(row=r_idx, column=17, value=t["deliverables"] or "Not Available")
        ws2.cell(row=r_idx, column=18, value=depends_on_str)
        ws2.cell(row=r_idx, column=19, value=blocks_str)

        for c in range(1, 20):
            ws2.cell(row=r_idx, column=c).border = thin_border
            if c != 3:
                ws2.cell(row=r_idx, column=c).font = regular_font
        r_idx += 1

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    # 3. Attention Required Sheet
    ws3 = wb.create_sheet(title="Attention Required")
    ws3.views.sheetView[0].showGridLines = True

    ws3["A1"] = "Critical Attention Required — Overdue & High Priority Items"
    ws3["A1"].font = section_font
    ws3["A1"].fill = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")
    ws3.merge_cells("A1:F1")

    risk_headers = ["Category", "Activity #", "Activity Title", "Assignee", "Due Date", "Status / Delay"]
    for col_idx, h in enumerate(risk_headers, 1):
        cell = ws3.cell(row=2, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = blue_header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    r_idx = 3
    for o in report.get("attention_required", {}).get("overdue", []):
        ws3.cell(row=r_idx, column=1, value="OVERDUE").fill = prio_urgent_fill
        ws3.cell(row=r_idx, column=2, value=f"#{o['seq_num']:02d}").alignment = Alignment(horizontal="center")
        ws3.cell(row=r_idx, column=3, value=o["title"]).font = bold_font
        ws3.cell(row=r_idx, column=4, value=o["assignee_name"])
        ws3.cell(row=r_idx, column=5, value=o["due_date"]).alignment = Alignment(horizontal="center")
        ws3.cell(row=r_idx, column=6, value=f"Overdue by {o.get('delay_days', 1)} day(s)").font = Font(color="DC2626", bold=True)
        for c in range(1, 7):
            ws3.cell(row=r_idx, column=c).border = thin_border
        r_idx += 1

    for u in report.get("attention_required", {}).get("urgent_high_pending", []):
        ws3.cell(row=r_idx, column=1, value=f"PENDING ({u['priority'].upper()})").fill = prio_high_fill
        ws3.cell(row=r_idx, column=2, value=f"#{u['seq_num']:02d}").alignment = Alignment(horizontal="center")
        ws3.cell(row=r_idx, column=3, value=u["title"]).font = bold_font
        ws3.cell(row=r_idx, column=4, value=u["assignee_name"])
        ws3.cell(row=r_idx, column=5, value=u["due_date"]).alignment = Alignment(horizontal="center")
        ws3.cell(row=r_idx, column=6, value=u["status"].upper())
        for c in range(1, 7):
            ws3.cell(row=r_idx, column=c).border = thin_border
        r_idx += 1

    for col in ws3.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws3.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # 4. Milestones Sheet
    ws4 = wb.create_sheet(title="Milestones Roadmap")
    ws4.views.sheetView[0].showGridLines = True

    ws4["A1"] = "Project Milestones & Deliverables Roadmap"
    ws4["A1"].font = section_font
    ws4["A1"].fill = accent_fill
    ws4.merge_cells("A1:D1")

    m_headers = ["Milestone ID", "Milestone Title", "Target Due Date", "Status"]
    for col_idx, h in enumerate(m_headers, 1):
        cell = ws4.cell(row=2, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = blue_header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    r_idx = 3
    for m in report.get("milestones", []):
        ws4.cell(row=r_idx, column=1, value=f"#{m['id']}").alignment = Alignment(horizontal="center")
        ws4.cell(row=r_idx, column=2, value=m["title"]).font = bold_font
        ws4.cell(row=r_idx, column=3, value=m["due_date"]).alignment = Alignment(horizontal="center")
        ws4.cell(row=r_idx, column=4, value=m["status"].upper()).alignment = Alignment(horizontal="center")
        for c in range(1, 5):
            ws4.cell(row=r_idx, column=c).border = thin_border
        r_idx += 1

    for col in ws4.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws4.column_dimensions[col_letter].width = max(max_len + 3, 16)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()

@app.get("/api/projects/<project_id:int>/cumulative-report")
def get_cumulative_report_endpoint(project_id):
    curr_user = get_current_user()
    with get_db() as conn:
        report_data = build_cumulative_project_report(conn, project_id, current_user=curr_user)
        if not report_data:
            return json_response({"error": "Project not found"}, status=404)
        return json_response(report_data)

@app.get("/api/projects/<project_id:int>/cumulative-report/export")
def export_cumulative_report_endpoint(project_id):
    curr_user = get_current_user()
    with get_db() as conn:
        report_data = build_cumulative_project_report(conn, project_id, current_user=curr_user)
        if not report_data:
            return json_response({"error": "Project not found"}, status=404)
        
        excel_bytes = generate_cumulative_project_excel(report_data)
        safe_name = re.sub(r'[^a-zA-Z0-9_\-]', '_', report_data['project']['name'])
        filename = f"{safe_name}_Cumulative_Activity_Report.xlsx"

        response.content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        response.set_header("Content-Disposition", f'attachment; filename="{filename}"')
        return excel_bytes

# ==================== EXPORT & IMPORT ====================

@app.get("/api/projects/<project_id:int>/export")
def export_project(project_id):
    with get_db() as conn:
        project = conn.execute("SELECT * FROM projects WHERE id = ?", (project_id,)).fetchone()
        if not project:
            return json_response({"error": "Project not found"}, status=404)
        
        members = [dict(m) for m in conn.execute("SELECT * FROM members WHERE project_id = ?", (project_id,)).fetchall()]
        sprints = [dict(s) for s in conn.execute("SELECT * FROM sprints WHERE project_id = ?", (project_id,)).fetchall()]
        milestones = [dict(m) for m in conn.execute("SELECT * FROM milestones WHERE project_id = ?", (project_id,)).fetchall()]
        tasks = conn.execute("SELECT * FROM tasks WHERE project_id = ?", (project_id,)).fetchall()

        all_tasks = []
        for t in tasks:
            t_dict = dict(t)
            t_dict["subtasks"] = [dict(s) for s in conn.execute("SELECT * FROM subtasks WHERE task_id = ?", (t["id"],)).fetchall()]
            t_dict["timelogs"] = [dict(tl) for tl in conn.execute("SELECT * FROM timelogs WHERE task_id = ?", (t["id"],)).fetchall()]
            all_tasks.append(t_dict)

        export_data = {
            "version": "1.0",
            "exported_at": get_now_iso(),
            "project": dict(project),
            "members": members,
            "sprints": sprints,
            "milestones": milestones,
            "tasks": all_tasks
        }
        return json_response(export_data)

@app.post("/api/projects/import")
def import_project():
    data = request.json or {}
    p_data = data.get("project", {})
    if not p_data or "name" not in p_data:
        return json_response({"error": "Invalid project import structure"}, status=400)

    now_str = get_now_iso()
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO projects (name, description, color, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?)
        """, (
            p_data.get("name") + " (Imported)",
            p_data.get("description", ""),
            p_data.get("color", "#3B82F6"),
            now_str,
            now_str
        ))
        new_p_id = cursor.lastrowid

        member_map = {}
        for m in data.get("members", []):
            cursor.execute("""
                INSERT INTO members (project_id, name, email, role, avatar_color)
                VALUES (?, ?, ?, ?, ?)
            """, (new_p_id, m["name"], m.get("email", ""), m.get("role", "Member"), m.get("avatar_color", "#6366F1")))
            member_map[m["id"]] = cursor.lastrowid

        sprint_map = {}
        for s in data.get("sprints", []):
            cursor.execute("""
                INSERT INTO sprints (project_id, name, goal, start_date, end_date, status)
                VALUES (?, ?, ?, ?, ?, ?)
            """, (new_p_id, s["name"], s.get("goal", ""), s.get("start_date"), s.get("end_date"), s.get("status", "planning")))
            sprint_map[s["id"]] = cursor.lastrowid

        for m in data.get("milestones", []):
            cursor.execute("""
                INSERT INTO milestones (project_id, title, due_date, status)
                VALUES (?, ?, ?, ?)
            """, (new_p_id, m["title"], m["due_date"], m.get("status", "pending")))

        for t in data.get("tasks", []):
            new_sprint_id = sprint_map.get(t.get("sprint_id"))
            new_assignee_id = member_map.get(t.get("assignee_id"))
            cursor.execute("""
                INSERT INTO tasks (
                    project_id, sprint_id, title, description, status, priority,
                    order_index, start_date, due_date, estimated_hours, actual_hours,
                    assignee_id, tags, created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                new_p_id, new_sprint_id, t["title"], t.get("description", ""),
                t.get("status", "todo"), t.get("priority", "medium"), t.get("order_index", 0),
                t.get("start_date"), t.get("due_date"), t.get("estimated_hours", 0.0),
                t.get("actual_hours", 0.0), new_assignee_id, json.dumps(t.get("tags", [])),
                now_str, now_str
            ))
            new_t_id = cursor.lastrowid

            for sub in t.get("subtasks", []):
                cursor.execute("""
                    INSERT INTO subtasks (task_id, title, completed, order_index)
                    VALUES (?, ?, ?, ?)
                """, (new_t_id, sub["title"], sub.get("completed", 0), sub.get("order_index", 0)))

            for tl in t.get("timelogs", []):
                new_m_id = member_map.get(tl.get("member_id"))
                cursor.execute("""
                    INSERT INTO timelogs (task_id, member_id, hours, description, logged_date, created_at)
                    VALUES (?, ?, ?, ?, ?, ?)
                """, (new_t_id, new_m_id, tl["hours"], tl.get("description", ""), tl.get("logged_date", now_str[:10]), now_str))

        record_activity(conn, new_p_id, "System", "Project Imported", "Imported project backup")
        bootstrap = get_bootstrap_payload(conn, active_project_id=new_p_id)
        return json_response({
            "success": True,
            "project_id": new_p_id,
            "projects": bootstrap["projects"],
            "current_project": bootstrap["current_project"],
            "tasks": bootstrap["tasks"]
        })

# ==================== GANTT EXCEL / CSV UPLOAD ====================

@app.get("/api/gantt/sample_csv")
def get_sample_gantt_csv():
    response.content_type = "text/csv; charset=utf-8"
    response.headers["Content-Disposition"] = "attachment; filename=\"sample_gantt_schedule.csv\""
    return generate_sample_gantt_csv()

@app.get("/api/gantt/sample_xlsx")
def get_sample_gantt_xlsx():
    response.content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    response.headers["Content-Disposition"] = "attachment; filename=\"sample_gantt_schedule.xlsx\""
    return generate_sample_gantt_excel()

@app.post("/api/projects/<project_id:int>/upload_gantt")
def upload_gantt_file(project_id):
    upload = request.files.get("file")
    if not upload:
        return json_response({"error": "No file uploaded"}, status=400)

    filename = upload.filename or "gantt.xlsx"
    file_bytes = upload.file.read()
    
    new_project_name = request.forms.get("new_project_name", "").strip()
    mode = request.forms.get("mode", "append")

    try:
        parsed = parse_gantt_file(file_bytes, filename)
    except Exception as err:
        return json_response({"error": f"Failed to parse Gantt file: {str(err)}"}, status=400)

    tasks_data = parsed["tasks"]
    if not tasks_data:
        return json_response({"error": "No valid task rows found in uploaded file"}, status=400)

    now_str = get_now_iso()

    with get_db() as conn:
        cursor = conn.cursor()
        
        target_p_id = project_id
        if new_project_name:
            cursor.execute("""
                INSERT INTO projects (name, description, color, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?)
            """, (new_project_name, f"Imported from {filename}", "#3B82F6", now_str, now_str))
            target_p_id = cursor.lastrowid
        else:
            project = conn.execute("SELECT * FROM projects WHERE id = ?", (project_id,)).fetchone()
            if not project:
                return json_response({"error": "Project not found"}, status=404)
            if mode == "replace":
                conn.execute("DELETE FROM tasks WHERE project_id = ?", (project_id,))

        # Match or create members
        existing_members = conn.execute("SELECT id, name FROM members WHERE project_id = ?", (target_p_id,)).fetchall()
        member_map = {m["name"].lower(): m["id"] for m in existing_members}
        new_members_count = 0

        for idx, m_name in enumerate(parsed["members_found"]):
            if m_name.lower() not in member_map:
                color = AVATAR_COLORS[idx % len(AVATAR_COLORS)]
                cursor.execute("""
                    INSERT INTO members (project_id, name, email, role, avatar_color)
                    VALUES (?, ?, ?, ?, ?)
                """, (target_p_id, m_name, f"{m_name.lower().replace(' ', '.')}@company.internal", "Contributor", color))
                new_m_id = cursor.lastrowid
                member_map[m_name.lower()] = new_m_id
                new_members_count += 1

        # Match or create sprints / phases
        existing_sprints = conn.execute("SELECT id, name FROM sprints WHERE project_id = ?", (target_p_id,)).fetchall()
        sprint_map = {s["name"].lower(): s["id"] for s in existing_sprints}
        new_sprints_count = 0

        for s_name in parsed["sprints_found"]:
            if s_name.lower() not in sprint_map:
                s_tasks = [t for t in tasks_data if t.get("sprint_name") == s_name]
                starts = [t["start_date"] for t in s_tasks if t.get("start_date")]
                ends = [t["due_date"] for t in s_tasks if t.get("due_date")]
                min_s = min(starts) if starts else now_str[:10]
                max_e = max(ends) if ends else (datetime.now() + timedelta(days=14)).strftime("%Y-%m-%d")

                cursor.execute("""
                    INSERT INTO sprints (project_id, name, goal, start_date, end_date, status)
                    VALUES (?, ?, ?, ?, ?, ?)
                """, (target_p_id, s_name, f"Phase goals for {s_name}", min_s, max_e, "active" if new_sprints_count == 0 else "planning"))
                new_s_id = cursor.lastrowid
                sprint_map[s_name.lower()] = new_s_id
                new_sprints_count += 1

        # Insert tasks
        max_order = conn.execute(
            "SELECT COALESCE(MAX(order_index), -1) as m FROM tasks WHERE project_id = ?",
            (target_p_id,)
        ).fetchone()["m"]

        task_rows = []
        for idx, t in enumerate(tasks_data):
            assignee_id = member_map.get(t["assignee_name"].lower()) if t.get("assignee_name") else None
            sprint_id = sprint_map.get(t["sprint_name"].lower()) if t.get("sprint_name") else None
            tags_json = json.dumps(t.get("tags") or [])
            task_rows.append((
                target_p_id, sprint_id, t["title"], t.get("description", ""),
                t.get("status", "todo"), t.get("priority", "medium"),
                max_order + idx + 1, t.get("start_date"), t.get("due_date"),
                float(t.get("estimated_hours") or 0.0), 0.0,
                assignee_id, tags_json, now_str, now_str
            ))

        cursor.executemany("""
            INSERT INTO tasks (
                project_id, sprint_id, title, description, status, priority,
                order_index, start_date, due_date, estimated_hours, actual_hours,
                assignee_id, tags, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, task_rows)

        record_activity(
            conn, target_p_id, "User", "Gantt Upload",
            f"Imported {len(tasks_data)} tasks with {len(parsed['members_found'])} assigned members from '{filename}'"
        )

        bootstrap = get_bootstrap_payload(conn, active_project_id=target_p_id)

        return json_response({
            "success": True,
            "tasks_imported": len(tasks_data),
            "members_added": new_members_count,
            "sprints_added": new_sprints_count,
            "project_id": target_p_id,
            "projects": bootstrap["projects"],
            "current_project": bootstrap["current_project"],
            "tasks": bootstrap["tasks"],
            "message": f"Successfully imported {len(tasks_data)} tasks with assigned members!"
        })

@app.post("/api/seed/reset")
def reset_database():
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("DROP TABLE IF EXISTS activity_logs")
        cursor.execute("DROP TABLE IF EXISTS timelogs")
        cursor.execute("DROP TABLE IF EXISTS subtasks")
        cursor.execute("DROP TABLE IF EXISTS tasks")
        cursor.execute("DROP TABLE IF EXISTS milestones")
        cursor.execute("DROP TABLE IF EXISTS sprints")
        cursor.execute("DROP TABLE IF EXISTS members")
        cursor.execute("DROP TABLE IF EXISTS projects")
        cursor.execute("DROP TABLE IF EXISTS notification_dispatches")
        cursor.execute("DROP TABLE IF EXISTS email_logs")
        cursor.execute("DROP TABLE IF EXISTS email_settings")
    init_db()
    seed_database()
    return json_response({"success": True, "message": "Database reset and demo data loaded"})

# ==================== NOTIFICATIONS & OUTLOOK INTEGRATION ====================

@app.get("/api/notifications/settings")
def get_notification_settings():
    with get_db() as conn:
        settings = get_settings(conn)
        masked = dict(settings)
        if masked.get("smtp_pass"):
            masked["has_password"] = True
            masked["smtp_pass"] = "••••••••"
        else:
            masked["has_password"] = False
        return json_response(masked)

@app.put("/api/notifications/settings")
def update_notification_settings():
    data = request.json or {}
    with get_db() as conn:
        current = get_settings(conn)
        
        provider = data.get("smtp_provider", current.get("smtp_provider", "outlook"))
        host = data.get("smtp_host", current.get("smtp_host", "smtp.office365.com"))
        port = int(data.get("smtp_port", current.get("smtp_port", 587)))
        user = data.get("smtp_user", current.get("smtp_user", ""))
        
        new_pass = data.get("smtp_pass")
        if new_pass is None or new_pass == "••••••••" or new_pass == "":
            password = current.get("smtp_pass", "")
        else:
            password = new_pass
            
        sender_name = data.get("sender_name", current.get("sender_name", "ProjectPulse Notifications"))
        use_tls = int(data.get("use_tls", current.get("use_tls", 1)))
        is_enabled = int(data.get("is_enabled", current.get("is_enabled", 1)))
        simulation_mode = int(data.get("simulation_mode", current.get("simulation_mode", 1)))
        notify_due_soon = int(data.get("notify_due_soon", current.get("notify_due_soon", 1)))
        notify_due_today = int(data.get("notify_due_today", current.get("notify_due_today", 1)))
        notify_overdue = int(data.get("notify_overdue", current.get("notify_overdue", 1)))
        recurring_day = data.get("recurring_day", current.get("recurring_day", "monday")).lower()
        notify_assigned = int(data.get("notify_assigned", current.get("notify_assigned", 1)))
        notify_completed = int(data.get("notify_completed", current.get("notify_completed", 1)))
        now_str = get_now_iso()

        conn.execute("""
            UPDATE email_settings SET
                smtp_provider = ?, smtp_host = ?, smtp_port = ?, smtp_user = ?,
                smtp_pass = ?, sender_name = ?, use_tls = ?, is_enabled = ?,
                simulation_mode = ?, notify_due_soon = ?, notify_due_today = ?,
                notify_overdue = ?, recurring_day = ?, notify_assigned = ?,
                notify_completed = ?, updated_at = ?
            WHERE id = 1
        """, (
            provider, host, port, user, password, sender_name, use_tls,
            is_enabled, simulation_mode, notify_due_soon, notify_due_today,
            notify_overdue, recurring_day, notify_assigned, notify_completed, now_str
        ))

        updated = get_settings(conn)
        masked = dict(updated)
        if masked.get("smtp_pass"):
            masked["has_password"] = True
            masked["smtp_pass"] = "••••••••"
        else:
            masked["has_password"] = False
        return json_response(masked)

@app.post("/api/notifications/test_email")
def test_email_endpoint():
    data = request.json or {}
    email = data.get("email", "").strip()
    if not email or "@" not in email:
        return json_response({"error": "Valid email address is required"}, status=400)

    with get_db() as conn:
        result = send_test_email(conn, email)
        return json_response(result)

@app.post("/api/notifications/run_checks")
def run_due_date_checks_endpoint():
    data = request.json or {}
    ref_date = data.get("reference_date")
    with get_db() as conn:
        dispatched = run_all_due_date_checks(conn, reference_date_str=ref_date)
        return json_response({
            "success": True,
            "checks_run_at": get_now_iso(),
            "notifications_dispatched": len(dispatched),
            "details": dispatched
        })

@app.get("/api/notifications/logs")
def get_notification_logs():
    limit = int(request.query.get("limit", 50))
    with get_db() as conn:
        logs = conn.execute("""
            SELECT l.id, l.project_id, l.task_id, l.recipient_email, l.recipient_name,
                   l.subject, l.trigger_type, l.status, l.error_message, l.sent_at,
                   p.name as project_name, t.title as task_title
            FROM email_logs l
            LEFT JOIN projects p ON l.project_id = p.id
            LEFT JOIN tasks t ON l.task_id = t.id
            ORDER BY l.sent_at DESC, l.id DESC
            LIMIT ?
        """, (limit,)).fetchall()
        return json_response(logs)

@app.get("/api/notifications/logs/<log_id:int>")
def get_notification_log_detail(log_id):
    with get_db() as conn:
        log = conn.execute("""
            SELECT l.*, p.name as project_name, t.title as task_title
            FROM email_logs l
            LEFT JOIN projects p ON l.project_id = p.id
            LEFT JOIN tasks t ON l.task_id = t.id
            WHERE l.id = ?
        """, (log_id,)).fetchone()
        if not log:
            return json_response({"error": "Log not found"}, status=404)
        return json_response(dict(log))

# ==================== RESOURCE MANAGEMENT & RESOURCE MAPPING ====================

def compute_resource_availability(total_alloc: float, status: str = "active") -> str:
    if status != "active":
        return "unavailable"
    if total_alloc <= 0:
        return "available"
    elif total_alloc < 100:
        return "partially_allocated"
    elif total_alloc == 100:
        return "fully_allocated"
    else:
        return "overallocated"

def get_resource_dict(conn, resource_id: int):
    r = conn.execute("SELECT * FROM resources WHERE id = ?", (resource_id,)).fetchone()
    if not r:
        return None
    
    r_dict = dict(r)
    try:
        r_dict["skills"] = json.loads(r_dict["skills"]) if r_dict.get("skills") else []
    except Exception:
        r_dict["skills"] = []

    # Total allocation across all active project mappings
    alloc_row = conn.execute("""
        SELECT COALESCE(SUM(allocation_pct), 0) as total_alloc
        FROM project_resources
        WHERE resource_id = ? AND status = 'active'
    """, (resource_id,)).fetchone()
    total_alloc = safe_float(alloc_row["total_alloc"] if alloc_row else 0.0)
    r_dict["total_allocation_pct"] = total_alloc
    r_dict["computed_availability_status"] = compute_resource_availability(total_alloc, r_dict.get("status", "active"))
    r_dict["computed_availability"] = r_dict["computed_availability_status"]
    r_dict["cost_per_hour"] = safe_float(r_dict.get("cost_rate", 0.0))
    r_dict["currency"] = r_dict.get("cost_unit", "USD")
    r_dict["availability_status"] = r_dict.get("availability", "available")
    r_dict["phone"] = r_dict.get("contact_phone", "")

    # Project mappings
    proj_mappings = conn.execute("""
        SELECT pr.*, pr.role as project_role, pr.allocation_pct as allocation_percentage,
               p.name as project_name, p.color as project_color
        FROM project_resources pr
        JOIN projects p ON pr.project_id = p.id
        WHERE pr.resource_id = ?
        ORDER BY pr.id DESC
    """, (resource_id,)).fetchall()
    r_dict["project_mappings"] = [dict(pm) for pm in proj_mappings]
    r_dict["project_mappings_count"] = len(r_dict["project_mappings"])
    r_dict["mapped_project_count"] = len(r_dict["project_mappings"])

    # Task assignments
    task_mappings = conn.execute("""
        SELECT tr.*, t.title as task_title, t.status as task_status, t.priority as task_priority,
               t.due_date as task_due_date, p.name as project_name, p.color as project_color
        FROM task_resources tr
        JOIN tasks t ON tr.task_id = t.id
        JOIN projects p ON tr.project_id = p.id
        WHERE tr.resource_id = ?
        ORDER BY tr.id DESC
    """, (resource_id,)).fetchall()
    r_dict["task_mappings"] = [dict(tm) for tm in task_mappings]
    r_dict["assigned_tasks_count"] = len(r_dict["task_mappings"])
    r_dict["assigned_task_count"] = len(r_dict["task_mappings"])
    r_dict["assigned_tasks"] = [
        {
            "id": tm["task_id"],
            "title": tm.get("task_title", ""),
            "project_name": tm.get("project_name", ""),
            "status": tm.get("task_status", ""),
            "priority": tm.get("task_priority", ""),
            "due_date": tm.get("task_due_date", "")
        }
        for tm in r_dict["task_mappings"]
    ]

    return r_dict

@app.get("/api/resources")
def get_resources():
    res_type = request.query.get("type")
    category = request.query.get("category")
    department = request.query.get("department")
    status = request.query.get("status")
    search = request.query.get("search")
    availability_filter = request.query.get("availability")

    with get_db() as conn:
        query = "SELECT * FROM resources WHERE 1=1"
        params = []

        if res_type:
            query += " AND type = ?"
            params.append(res_type)
        if category:
            query += " AND category = ?"
            params.append(category)
        if department:
            query += " AND department = ?"
            params.append(department)
        if status:
            query += " AND status = ?"
            params.append(status)
        if search:
            query += " AND (name LIKE ? OR resource_code LIKE ? OR role LIKE ? OR skills LIKE ? OR department LIKE ?)"
            s_param = f"%{search}%"
            params.extend([s_param, s_param, s_param, s_param, s_param])

        query += " ORDER BY name ASC"
        rows = conn.execute(query, params).fetchall()

        # Batch compute allocation & counts
        alloc_rows = conn.execute("""
            SELECT resource_id, COALESCE(SUM(allocation_pct), 0) as total_alloc, COUNT(DISTINCT project_id) as proj_cnt
            FROM project_resources
            WHERE status = 'active'
            GROUP BY resource_id
        """).fetchall()
        alloc_map = {row["resource_id"]: (safe_float(row["total_alloc"]), row["proj_cnt"]) for row in alloc_rows}

        task_cnt_rows = conn.execute("""
            SELECT resource_id, COUNT(DISTINCT task_id) as task_cnt
            FROM task_resources
            GROUP BY resource_id
        """).fetchall()
        task_cnt_map = {row["resource_id"]: row["task_cnt"] for row in task_cnt_rows}

        results = []
        for r in rows:
            r_dict = dict(r)
            try:
                r_dict["skills"] = json.loads(r_dict["skills"]) if r_dict.get("skills") else []
            except Exception:
                r_dict["skills"] = []

            total_alloc, proj_cnt = alloc_map.get(r["id"], (0.0, 0))
            r_dict["total_allocation_pct"] = total_alloc
            comp_avail = compute_resource_availability(total_alloc, r_dict.get("status", "active"))
            r_dict["computed_availability_status"] = comp_avail
            r_dict["computed_availability"] = comp_avail
            r_dict["cost_per_hour"] = safe_float(r_dict.get("cost_rate", 0.0))
            r_dict["currency"] = r_dict.get("cost_unit", "USD")
            r_dict["availability_status"] = r_dict.get("availability", "available")
            r_dict["phone"] = r_dict.get("contact_phone", "")
            r_dict["project_mappings_count"] = proj_cnt
            r_dict["mapped_project_count"] = proj_cnt
            r_dict["assigned_tasks_count"] = task_cnt_map.get(r["id"], 0)
            r_dict["assigned_task_count"] = task_cnt_map.get(r["id"], 0)

            if availability_filter and comp_avail != availability_filter:
                continue

            results.append(r_dict)

        return json_response(results)

@app.post("/api/resources")
def create_resource():
    data = request.json or {}
    name = clean_text(data.get("name"))
    res_type = clean_text(data.get("type") or "Employee")
    
    if not name:
        return json_response({"error": "Resource name is required"}, status=400)

    category = clean_text(data.get("category") or "Internal")
    department = clean_text(data.get("department") or "Engineering")
    role = clean_text(data.get("role") or "")
    status = clean_text(data.get("status") or "active")
    availability = clean_text(data.get("availability_status") or data.get("availability") or "available")
    description = clean_text(data.get("description") or "")
    contact_email = clean_text(data.get("contact_email") or "")
    contact_phone = clean_text(data.get("contact_phone") or data.get("phone") or "")
    location = clean_text(data.get("location") or "")
    cost_rate = safe_float(data.get("cost_per_hour") if data.get("cost_per_hour") is not None else data.get("cost_rate"), 0.0)
    cost_unit = clean_text(data.get("currency") or data.get("cost_unit") or "USD")
    notes = clean_text(data.get("notes") or "")

    raw_skills = data.get("skills", [])
    if isinstance(raw_skills, list):
        skills_list = [str(s).strip() for s in raw_skills if str(s).strip()]
    elif isinstance(raw_skills, str):
        skills_list = [s.strip() for s in raw_skills.split(',') if s.strip()]
    else:
        skills_list = []
    skills_json = json.dumps(skills_list)

    now_str = get_now_iso()

    with get_db() as conn:
        cursor = conn.cursor()
        
        # Auto-generate resource code if not provided
        resource_code = clean_text(data.get("resource_code"))
        if not resource_code:
            prefix_map = {
                "Employee": "EMP",
                "Contractor": "CTR",
                "Equipment": "EQP",
                "Laboratory Equipment": "LAB",
                "Software": "SFT",
                "Vendor": "VND",
                "External Resource": "EXT",
                "Material": "MAT",
                "Facility": "FAC"
            }
            pfx = prefix_map.get(res_type, "RES")
            count_row = conn.execute("SELECT COUNT(*) as cnt FROM resources WHERE type = ?", (res_type,)).fetchone()
            next_num = (count_row["cnt"] if count_row else 0) + 101
            resource_code = f"{pfx}-{next_num:03d}"
            
            # Ensure unique
            dup = conn.execute("SELECT id FROM resources WHERE resource_code = ?", (resource_code,)).fetchone()
            if dup:
                import random
                resource_code = f"{pfx}-{random.randint(1000, 9999)}"

        try:
            cursor.execute("""
                INSERT INTO resources (
                    resource_code, name, type, category, department, role, skills,
                    description, availability, status, contact_email, contact_phone, location,
                    cost_rate, cost_unit, notes, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                resource_code, name, res_type, category, department, role, skills_json,
                description, availability, status, contact_email, contact_phone, location,
                cost_rate, cost_unit, notes, now_str, now_str
            ))
            r_id = cursor.lastrowid
            
            # Record activity on first project if available
            p_first = conn.execute("SELECT id FROM projects ORDER BY id ASC LIMIT 1").fetchone()
            if p_first:
                record_activity(conn, p_first["id"], "Manager", "Resource Created", f'Created resource "{name}" ({resource_code})')

            return json_response(get_resource_dict(conn, r_id), status=201)
        except sqlite3.IntegrityError as e:
            return json_response({"error": f"Resource code or name already exists: {e}"}, status=400)

@app.get("/api/resources/summary")
def get_resources_summary():
    with get_db() as conn:
        all_resources = conn.execute("SELECT * FROM resources").fetchall()
        total_count = len(all_resources)
        active_count = sum(1 for r in all_resources if r["status"] == "active")

        # Allocations
        alloc_rows = conn.execute("""
            SELECT resource_id, COALESCE(SUM(allocation_pct), 0) as total_alloc
            FROM project_resources
            WHERE status = 'active'
            GROUP BY resource_id
        """).fetchall()
        alloc_map = {row["resource_id"]: safe_float(row["total_alloc"]) for row in alloc_rows}

        by_type = {}
        by_dept = {}
        by_avail = {
            "available": 0,
            "partially_allocated": 0,
            "fully_allocated": 0,
            "overallocated": 0,
            "unavailable": 0
        }

        total_alloc_sum = 0.0
        for r in all_resources:
            t = r["type"] or "Other"
            by_type[t] = by_type.get(t, 0) + 1

            d = r["department"] or "General"
            by_dept[d] = by_dept.get(d, 0) + 1

            alloc = alloc_map.get(r["id"], 0.0)
            total_alloc_sum += alloc
            avail = compute_resource_availability(alloc, r.get("status", "active"))
            by_avail[avail] = by_avail.get(avail, 0) + 1

        avg_alloc = (total_alloc_sum / total_count) if total_count > 0 else 0.0

        proj_res_count = conn.execute("SELECT COUNT(*) as cnt FROM project_resources").fetchone()["cnt"]
        task_res_count = conn.execute("SELECT COUNT(*) as cnt FROM task_resources").fetchone()["cnt"]

        return json_response({
            "total_resources": total_count,
            "active_resources": active_count,
            "avg_allocation_pct": round(avg_alloc, 1),
            "average_allocation_pct": round(avg_alloc, 1),
            "total_project_mappings": proj_res_count,
            "total_assigned_tasks": task_res_count,
            "total_task_assignments": task_res_count,
            "by_type": by_type,
            "by_department": by_dept,
            "by_availability": by_avail
        })

@app.get("/api/resources/<resource_id:int>")
def get_resource_detail(resource_id):
    with get_db() as conn:
        r_dict = get_resource_dict(conn, resource_id)
        if not r_dict:
            return json_response({"error": "Resource not found"}, status=404)
        return json_response(r_dict)

@app.put("/api/resources/<resource_id:int>")
def update_resource(resource_id):
    data = request.json or {}
    with get_db() as conn:
        existing = conn.execute("SELECT * FROM resources WHERE id = ?", (resource_id,)).fetchone()
        if not existing:
            return json_response({"error": "Resource not found"}, status=404)

        name = clean_text(data.get("name", existing["name"]))
        resource_code = clean_text(data.get("resource_code", existing["resource_code"]))
        res_type = clean_text(data.get("type", existing["type"]))
        category = clean_text(data.get("category", existing["category"]))
        department = clean_text(data.get("department", existing["department"]))
        role = clean_text(data.get("role", existing["role"]))
        status = clean_text(data.get("status", existing["status"]))
        availability = clean_text(data.get("availability_status", data.get("availability", existing["availability"])))
        description = clean_text(data.get("description", existing["description"]))
        contact_email = clean_text(data.get("contact_email", existing["contact_email"]))
        contact_phone = clean_text(data.get("contact_phone", data.get("phone", existing["contact_phone"])))
        location = clean_text(data.get("location", existing["location"]))
        cost_rate = safe_float(data.get("cost_per_hour") if data.get("cost_per_hour") is not None else data.get("cost_rate", existing["cost_rate"]))
        cost_unit = clean_text(data.get("currency", data.get("cost_unit", existing["cost_unit"])))
        notes = clean_text(data.get("notes", existing["notes"]))

        if "skills" in data:
            raw_skills = data["skills"]
            if isinstance(raw_skills, list):
                skills_list = [str(s).strip() for s in raw_skills if str(s).strip()]
            elif isinstance(raw_skills, str):
                skills_list = [s.strip() for s in raw_skills.split(',') if s.strip()]
            else:
                skills_list = []
            skills_json = json.dumps(skills_list)
        else:
            skills_json = existing["skills"] or "[]"

        now_str = get_now_iso()

        conn.execute("""
            UPDATE resources SET
                resource_code = ?, name = ?, type = ?, category = ?, department = ?,
                role = ?, skills = ?, description = ?, availability = ?, status = ?,
                contact_email = ?, contact_phone = ?, location = ?, cost_rate = ?,
                cost_unit = ?, notes = ?, updated_at = ?
            WHERE id = ?
        """, (
            resource_code, name, res_type, category, department,
            role, skills_json, description, availability, status,
            contact_email, contact_phone, location, cost_rate,
            cost_unit, notes, now_str, resource_id
        ))

        return json_response(get_resource_dict(conn, resource_id))

@app.delete("/api/resources/<resource_id:int>")
def delete_resource(resource_id):
    with get_db() as conn:
        conn.execute("DELETE FROM task_resources WHERE resource_id = ?", (resource_id,))
        conn.execute("DELETE FROM project_resources WHERE resource_id = ?", (resource_id,))
        conn.execute("DELETE FROM resources WHERE id = ?", (resource_id,))
        return json_response({"success": True})

# ==================== PROJECT RESOURCE MAPPING ====================

@app.get("/api/projects/<project_id:int>/resources")
def get_project_resources(project_id):
    with get_db() as conn:
        mappings = conn.execute("""
            SELECT pr.id as mapping_id, r.id as id, pr.project_id, pr.resource_id, pr.role, pr.role as project_role,
                   pr.allocation_pct, pr.allocation_pct as allocation_percentage,
                   pr.start_date, pr.start_date as mapping_start_date,
                   pr.end_date, pr.end_date as mapping_end_date,
                   pr.responsibility, pr.status, pr.status as mapping_status,
                   pr.created_at as mapping_created_at,
                   r.resource_code, r.name, r.type, r.category, r.department,
                   r.role as primary_role, r.skills, r.availability as availability_status,
                   r.status as resource_status, r.contact_email, r.contact_phone, r.contact_phone as phone,
                   r.location, r.cost_rate as cost_per_hour, r.cost_unit as currency
            FROM project_resources pr
            JOIN resources r ON pr.resource_id = r.id
            WHERE pr.project_id = ?
            ORDER BY r.name ASC
        """, (project_id,)).fetchall()

        # Get global allocation for each resource across ALL active projects
        alloc_rows = conn.execute("""
            SELECT resource_id, COALESCE(SUM(allocation_pct), 0) as total_alloc
            FROM project_resources
            WHERE status = 'active'
            GROUP BY resource_id
        """).fetchall()
        alloc_map = {row["resource_id"]: safe_float(row["total_alloc"]) for row in alloc_rows}

        # Count assigned tasks in this project
        task_cnt_rows = conn.execute("""
            SELECT resource_id, COUNT(DISTINCT task_id) as task_cnt
            FROM task_resources
            WHERE project_id = ?
            GROUP BY resource_id
        """, (project_id,)).fetchall()
        task_cnt_map = {row["resource_id"]: row["task_cnt"] for row in task_cnt_rows}

        results = []
        for m in mappings:
            m_dict = dict(m)
            try:
                m_dict["skills"] = json.loads(m_dict["skills"]) if m_dict.get("skills") else []
            except Exception:
                m_dict["skills"] = []
            
            total_alloc = alloc_map.get(m["resource_id"], 0.0)
            m_dict["total_allocation_pct"] = total_alloc
            m_dict["global_allocation_pct"] = total_alloc
            m_dict["computed_availability_status"] = compute_resource_availability(total_alloc, m_dict.get("resource_status", "active"))
            m_dict["computed_availability"] = m_dict["computed_availability_status"]
            m_dict["project_tasks_count"] = task_cnt_map.get(m["resource_id"], 0)
            results.append(m_dict)

        return json_response(results)

@app.post("/api/projects/<project_id:int>/resources")
def map_project_resource(project_id):
    data = request.json or {}
    resource_id = safe_int(data.get("resource_id"))
    if not resource_id:
        return json_response({"error": "Resource ID is required"}, status=400)

    project_role = clean_text(data.get("project_role") or data.get("role") or "")
    allocation_pct = safe_float(data.get("allocation_pct") if data.get("allocation_pct") is not None else data.get("allocation_percentage"), 100.0)
    start_date = data.get("start_date")
    start_date = str(start_date).strip() if start_date and str(start_date).strip() not in ("", "null", "undefined", "None") else None
    end_date = data.get("end_date")
    end_date = str(end_date).strip() if end_date and str(end_date).strip() not in ("", "null", "undefined", "None") else None
    responsibility = clean_text(data.get("responsibility") or "")
    status = clean_text(data.get("status") or "active")
    now_str = get_now_iso()

    with get_db() as conn:
        # Check resource exists
        r = conn.execute("SELECT * FROM resources WHERE id = ?", (resource_id,)).fetchone()
        if not r:
            return json_response({"error": "Resource not found"}, status=404)

        if not project_role:
            project_role = r["role"] or r["type"]

        cursor = conn.cursor()
        existing_map = conn.execute(
            "SELECT id FROM project_resources WHERE project_id = ? AND resource_id = ?",
            (project_id, resource_id)
        ).fetchone()

        if existing_map:
            cursor.execute("""
                UPDATE project_resources SET
                    role = ?, allocation_pct = ?, start_date = ?,
                    end_date = ?, responsibility = ?, status = ?, updated_at = ?
                WHERE id = ?
            """, (project_role, allocation_pct, start_date, end_date, responsibility, status, now_str, existing_map["id"]))
            mapping_id = existing_map["id"]
        else:
            cursor.execute("""
                INSERT INTO project_resources (
                    project_id, resource_id, role, allocation_pct,
                    start_date, end_date, responsibility, status, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (project_id, resource_id, project_role, allocation_pct, start_date, end_date, responsibility, status, now_str, now_str))
            mapping_id = cursor.lastrowid

        record_activity(conn, project_id, "Manager", "Resource Mapped", f'Mapped resource "{r["name"]}" as {project_role} ({allocation_pct}%)')

        mapping = conn.execute("""
            SELECT pr.id, pr.id as mapping_id, pr.project_id, pr.resource_id, pr.role, pr.role as project_role,
                   pr.allocation_pct, pr.allocation_pct as allocation_percentage,
                   pr.start_date, pr.start_date as mapping_start_date,
                   pr.end_date, pr.end_date as mapping_end_date,
                   pr.responsibility, pr.status, pr.status as mapping_status,
                   r.resource_code, r.name, r.type, r.department
            FROM project_resources pr
            JOIN resources r ON pr.resource_id = r.id
            WHERE pr.id = ?
        """, (mapping_id,)).fetchone()

        return json_response(dict(mapping), status=201)

@app.put("/api/projects/<project_id:int>/resources/<mapping_id:int>")
def update_project_resource_mapping(project_id, mapping_id):
    data = request.json or {}
    with get_db() as conn:
        existing = conn.execute("SELECT * FROM project_resources WHERE id = ? AND project_id = ?", (mapping_id, project_id)).fetchone()
        if not existing:
            return json_response({"error": "Project resource mapping not found"}, status=404)

        project_role = clean_text(data.get("project_role", data.get("role", existing["role"])))
        allocation_pct = safe_float(data.get("allocation_pct", data.get("allocation_percentage", existing["allocation_pct"])))
        start_date = data.get("start_date", existing["start_date"])
        start_date = str(start_date).strip() if start_date and str(start_date).strip() not in ("", "null", "undefined", "None") else None
        end_date = data.get("end_date", existing["end_date"])
        end_date = str(end_date).strip() if end_date and str(end_date).strip() not in ("", "null", "undefined", "None") else None
        responsibility = clean_text(data.get("responsibility", existing["responsibility"]))
        status = clean_text(data.get("status", existing["status"]))
        now_str = get_now_iso()

        conn.execute("""
            UPDATE project_resources SET
                role = ?, allocation_pct = ?, start_date = ?,
                end_date = ?, responsibility = ?, status = ?, updated_at = ?
            WHERE id = ?
        """, (project_role, allocation_pct, start_date, end_date, responsibility, status, now_str, mapping_id))

        mapping = conn.execute("""
            SELECT pr.id, pr.id as mapping_id, pr.project_id, pr.resource_id, pr.role, pr.role as project_role,
                   pr.allocation_pct, pr.allocation_pct as allocation_percentage,
                   pr.start_date, pr.start_date as mapping_start_date,
                   pr.end_date, pr.end_date as mapping_end_date,
                   pr.responsibility, pr.status, pr.status as mapping_status,
                   r.resource_code, r.name, r.type, r.department
            FROM project_resources pr
            JOIN resources r ON pr.resource_id = r.id
            WHERE pr.id = ?
        """, (mapping_id,)).fetchone()

        return json_response(dict(mapping))

@app.delete("/api/projects/<project_id:int>/resources/<resource_or_mapping_id:int>")
def unmap_project_resource(project_id, resource_or_mapping_id):
    with get_db() as conn:
        mapping = conn.execute(
            "SELECT * FROM project_resources WHERE (id = ? OR resource_id = ?) AND project_id = ?",
            (resource_or_mapping_id, resource_or_mapping_id, project_id)
        ).fetchone()
        if not mapping:
            return json_response({"error": "Project resource mapping not found"}, status=404)

        r = conn.execute("SELECT name FROM resources WHERE id = ?", (mapping["resource_id"],)).fetchone()
        r_name = r["name"] if r else "Resource"

        # Remove task resource assignments for this resource within this project
        conn.execute("DELETE FROM task_resources WHERE resource_id = ? AND project_id = ?", (mapping["resource_id"], project_id))
        conn.execute("DELETE FROM project_resources WHERE id = ?", (mapping["id"],))

        record_activity(conn, project_id, "Manager", "Resource Unmapped", f'Removed resource "{r_name}" from project')
        return json_response({"success": True})

# ==================== TASK RESOURCE MAPPING ====================

@app.get("/api/tasks/<task_id:int>/resources")
def get_task_resources(task_id):
    with get_db() as conn:
        task = conn.execute("SELECT project_id FROM tasks WHERE id = ?", (task_id,)).fetchone()
        if not task:
            return json_response({"error": "Task not found"}, status=404)

        trs = conn.execute("""
            SELECT tr.*, r.id as resource_id, r.id, r.resource_code, r.name as resource_name, r.name,
                   r.type as resource_type, r.type, r.category as resource_category, r.department as resource_department
            FROM task_resources tr
            JOIN resources r ON tr.resource_id = r.id
            WHERE tr.task_id = ?
            ORDER BY r.name ASC
        """, (task_id,)).fetchall()
        return json_response([dict(t) for t in trs])

@app.post("/api/tasks/<task_id:int>/resources")
def map_task_resource(task_id):
    data = request.json or {}
    resource_id = safe_int(data.get("resource_id"))
    if not resource_id:
        return json_response({"error": "Resource ID is required"}, status=400)

    with get_db() as conn:
        task = conn.execute("SELECT project_id, title FROM tasks WHERE id = ?", (task_id,)).fetchone()
        if not task:
            return json_response({"error": "Task not found"}, status=404)

        role = clean_text(data.get("role") or "")
        responsibility = clean_text(data.get("responsibility") or "")
        now_str = get_now_iso()

        conn.execute("""
            INSERT OR REPLACE INTO task_resources (task_id, resource_id, project_id, role, responsibility, created_at)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (task_id, resource_id, task["project_id"], role, responsibility, now_str))

        record_activity(conn, task["project_id"], "User", "Resource Assigned", f'Assigned resource to task "{task["title"]}"', task_id=task_id)
        return json_response({"success": True})

@app.delete("/api/tasks/<task_id:int>/resources/<resource_id:int>")
def unmap_task_resource(task_id, resource_id):
    with get_db() as conn:
        conn.execute("DELETE FROM task_resources WHERE task_id = ? AND resource_id = ?", (task_id, resource_id))
        return json_response({"success": True})

# ==================== GLOBAL JSON ERROR HANDLERS ====================

@app.error(400)
def error_400(error):
    response.content_type = "application/json"
    msg = str(error.body) if (error.body and not str(error.body).startswith("<!DOCTYPE")) else "Bad Request"
    return json.dumps({"error": msg, "status": 400})

@app.error(404)
def error_404(error):
    response.content_type = "application/json"
    msg = str(error.body) if (error.body and not str(error.body).startswith("<!DOCTYPE")) else "Resource not found"
    return json.dumps({"error": msg, "status": 404})

@app.error(405)
def error_405(error):
    response.content_type = "application/json"
    return json.dumps({"error": "Method Not Allowed", "status": 405})

@app.error(500)
def error_500(error):
    response.content_type = "application/json"
    if getattr(error, 'exception', None):
        import traceback
        traceback.print_exception(error.exception)
    msg = str(error.exception) if getattr(error, 'exception', None) else (str(error.body) if (error.body and not str(error.body).startswith("<!DOCTYPE")) else "Internal Server Error")
    return json.dumps({"error": msg, "status": 500})

def init_app():
    try:
        init_db()
        seed_database()
        start_background_scheduler()
    except Exception as e:
        print(f"[ProjectPulse Warning] Background initialization non-fatal error: {e}")

# Auto-initialize on module load (WSGI / Gunicorn / Bottle)
_initialized = False
if not _initialized:
    init_app()
    _initialized = True

def main():
    port = int(os.environ.get("PORT", 8080))
    host = os.environ.get("HOST", "0.0.0.0")
    print(f"============================================================")
    print(f"  ProjectPulse Production Server (Multi-threaded)")
    print(f"  Listening on: http://{host}:{port}")
    print(f"  -> Local access: http://localhost:{port}")
    print(f"============================================================")
    try:
        from waitress import serve
        serve(app, host=host, port=port, threads=8)
    except ImportError:
        run(app, host=host, port=port, reloader=False)

if __name__ == "__main__":
    main()

